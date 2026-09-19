"""Build 01 CSV, 02/03/04/08 workbooks from research/*.json. Values only (no formulas except summary counts)."""
import json, csv, re, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

R = "research"; O = "outputs"; DATE = "2026-09-02"
FONT = "Arial"
H = Font(name=FONT, bold=True, color="FFFFFF", size=10); HF = PatternFill("solid", fgColor="1F3D33")
B = Font(name=FONT, size=10); BOLD = Font(name=FONT, bold=True, size=10); T = Font(name=FONT, bold=True, size=14)
NOTE = Font(name=FONT, italic=True, size=9, color="444444")
LIMIT = ("RESEARCH LIMITATION (material): this audit was performed on 2026-09-02 from a research environment whose egress proxy blocked direct page loads "
         "of company websites, regulators, review sites, ad libraries and traffic tools, and whose web-search budget was capped. Facts therefore derive from "
         "dated search-result excerpts of the cited pages, labeled Verified fact / Company-reported claim / Third-party estimate / Analyst inference. "
         "Website UI/UX scores are either 'not assessed' or low-confidence inferences and must be re-scored from rendered pages before external use. "
         "Traffic estimates, social audience counts and ad-library activity were not retrievable. No screenshots were captured. Nothing was invented; 'unknown' means not found.")

def S(v):
    if v is None: return ""
    if isinstance(v, (list, tuple)): return "; ".join(S(x) for x in v)
    if isinstance(v, dict): return "; ".join(f"{k}: {S(x)}" for k, x in v.items())
    return str(v)

def sheet(wb, title, headers, rows, widths=None, note=None, freeze=True):
    ws = wb.create_sheet(title[:31])
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = NOTE
        ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(headers), 12))
        ws.row_dimensions[1].height = 75
        r0 = 3
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=r0, column=i, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
    for j, row in enumerate(rows, r0 + 1):
        for i, v in enumerate(row, 1):
            c = ws.cell(row=j, column=i, value=(v if isinstance(v, (int, float)) and not isinstance(v, bool) else S(v)))
            c.font = B; c.alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[L(i)].width = (widths[i - 1] if widths and i - 1 < len(widths) else 22)
    if freeze: ws.freeze_panes = ws.cell(row=r0 + 1, column=2)
    ws.auto_filter.ref = f"A{r0}:{L(len(headers))}{r0 + max(len(rows),1)}"
    return ws

def readme(wb, title, lines):
    ws = wb.active; ws.title = "README"
    ws["A1"] = title; ws["A1"].font = T
    for i, t in enumerate(lines, 3):
        c = ws.cell(row=i, column=1, value=t); c.font = B; c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = max(18, 15 * (len(t) // 130 + 1))
    ws.column_dimensions["A"].width = 150

d1 = [x for x in json.load(open(f"{R}/01_telehealth_companies.json")) if "company" in x]
for _c in d1:
    _c["segment_detail"] = _c.get("segment", ""); _c["segment"] = "Telehealth/clinician-led"
meta1 = [x for x in json.load(open(f"{R}/01_telehealth_companies.json")) if "_meta" in x][0]["_meta"]
d2 = json.load(open(f"{R}/02_pharmacy_ruo_consumer.json"))
d3 = json.load(open(f"{R}/03_celebrity_database.json"))
d4 = json.load(open(f"{R}/04_principal_verification.json"))
d5 = json.load(open(f"{R}/05_compliance_matrix.json"))
d6m = json.load(open(f"{R}/06_market_size.json"))
d6b = json.load(open(f"{R}/06_benchmarks.json"))

# ---------------- rankings (from research md files; within-segment)
RANKS = {
 "Telehealth/clinician-led": ["Hims", "Function Health", "Superpower", "Hone", "Joi", "Mochi", "Ivím", "Noom", "Peter MD", "Maximus", "Lifeforce", "Marek", "Defy", "Ro (", "Eden", "Henry", "TRT Nation", "Gameday", "AlphaMD", "Calibrate", "Limitless Male", "BodyLogic", "Fountain", "Concierge", "Mito", "Found", "Craft"],
 "Compounding pharmacy": ["Belmar", "Wells", "Olympia", "Red Rock", "Valor", "Innovation", "Foothills", "Hallandale", "Strive", "Empower", "ReviveRx", "Boothwyn", "ProRx", "Tailor Made"],
 "RUO vendor": ["Particle", "BioLongevity", "Polaris", "Core Peptides", "Biotech Peptides", "Sports Technology", "Aminos", "Limitless Life", "PureRawz", "Swiss Chems", "Peptide Sciences", "Amino Asylum", "Paradigm"],
 "Consumer supplement/cosmetic": ["Thorne", "Vital Proteins", "Medik8", "The Ordinary", "Momentous", "Timeline", "ProLon", "Ancient", "Drunk Elephant", "Olay", "BodyHealth", "Oral 'BPC-157'"],
 "Education/testing/practitioner platform": ["Function Health", "Superpower", "SSRP", "InsideTracker", "SiPhox", "A4M", "International Peptide", "Jay Campbell", "Hunter Williams"],
}
RATIONALE = {
 "Telehealth/clinician-led": "Proxy scorecard weighted to disclosed funding/performance, formulary breadth, pricing evidence and review signals (traffic/ads/social not retrievable). See research/01_telehealth_companies.md.",
 "Compounding pharmacy": "Ranked on enforcement record 2018-2026, verified 503A/503B footprint and testing infrastructure, peptide/GLP-1 menu breadth, resilience to 2026 rule changes. See research/02_pharmacy_ruo_consumer.md.",
 "RUO vendor": "Ranked on absence/presence of FDA action, published quality/testing claims, payment posture and marketing restraint; prices not retrievable. Segment is legally distinct: FDA treats human-use marketing as unapproved-drug distribution regardless of RUO labels.",
 "Consumer supplement/cosmetic": "Ranked on verified scale/ownership, strength of peptide positioning, substantiation/trust signals and 2026 momentum. Supplements (DSHEA) and cosmetics only; not drugs.",
 "Education/testing/practitioner platform": "Ranked on verified funding/scale, pricing transparency, relevance to peptide users/prescribers and regulatory posture.",
}
def rank_of(seg, name):
    lst = RANKS.get(seg, [])
    for i, k in enumerate(lst, 1):
        if k.lower() in name.lower(): return i
    return ""

def price_band(prods):
    ps = [S(p.get("displayed_price")) for p in prods if S(p.get("displayed_price")) not in ("", "unknown", "not retrieved", "Not verified this session", "not yet published")]
    return "; ".join(sorted(set(ps)))[:400]

def first_url(c):
    src = c.get("sources") or []
    return src[0]["url"] if src and isinstance(src[0], dict) else c.get("url", "")

def conf(c):
    lab = [s.get("label", "") for s in (c.get("sources") or []) if isinstance(s, dict)]
    if not lab: return "Low"
    v = sum(1 for l in lab if l.startswith("Verified")); cr = sum(1 for l in lab if l.startswith("Company"))
    return "Medium-High" if v >= 3 else ("Medium" if v + cr >= 2 else "Low")

# ---------------- 01 CSV
LEGAL = {
 "Telehealth/clinician-led": "Prescription telehealth via licensed clinicians; compounded or approved drugs dispensed by pharmacies",
 "Compounding pharmacy": "503A/503B pharmacy; dispenses only on prescription",
 "RUO vendor": "Sells 'research use only' chemicals; FDA treats human-use marketing as unapproved drug distribution",
 "Consumer supplement/cosmetic": "DSHEA supplements / cosmetics; structure-function or appearance claims only",
 "Education/testing/practitioner platform": "Education, diagnostics memberships or practitioner training; some route to Rx via partners",
 "Longevity/hormone clinic": "Clinic-based longevity/hormone programs; prescriptions via physicians",
 "Adjacent wellness": "Conventional wellness / fitness / nutrition; no drug claims",
 "Education/testing platform": "Diagnostics or testing memberships",
}
cols = ["company", "url", "segment", "legal_model", "business_model", "founded", "hq", "founders_leadership", "target_market", "geography_availability", "partners", "celebrity_involvement", "funding_or_disclosed_performance", "customer_promise", "positioning", "key_products", "displayed_price_examples", "traffic_proxy", "social_audience", "review_signal", "advertising_visibility", "regulatory_or_risk_notes", "within_segment_rank", "rank_methodology", "evidence_confidence", "primary_source_url", "access_date", "research_file"]
rows01 = []
for c in d1:
    rows01.append([c["company"], c.get("url"), c["segment"], LEGAL[c["segment"]], c.get("business_model"), c.get("founded"), c.get("hq"), c.get("founders_leadership"), c.get("target_market"), c.get("geography_availability"), c.get("clinical_pharmacy_partners"), c.get("celebrity_involvement"), c.get("funding_or_disclosed_performance"), c.get("customer_promise"), c.get("positioning"), "; ".join(S(p.get("product_or_program")) for p in c.get("products", []))[:400], price_band(c.get("products", [])), c.get("traffic_proxy"), c.get("social_audience"), c.get("review_signal"), c.get("advertising_visibility"), S(c.get("messaging_claims", {}).get("risky_claims") if isinstance(c.get("messaging_claims"), dict) else c.get("messaging_claims")), rank_of(c["segment"], c["company"]), RATIONALE[c["segment"]], conf(c), first_url(c), DATE, "research/01_telehealth_companies.json"])
for c in d2:
    rows01.append([c["company"], c.get("url"), c["segment"], LEGAL[c["segment"]], c.get("business_model"), c.get("founded"), c.get("hq"), c.get("founders_leadership"), c.get("target_market"), c.get("geography_availability"), c.get("partners"), c.get("celebrity_involvement"), c.get("funding_or_disclosed_performance"), c.get("customer_promise"), c.get("positioning"), "; ".join(S(p.get("product_or_program")) for p in c.get("products", []))[:400], price_band(c.get("products", [])), c.get("traffic_proxy"), c.get("social_audience"), c.get("review_signal"), c.get("advertising_visibility"), S(c.get("regulatory_notes"))[:600], rank_of(c["segment"], c["company"]), RATIONALE[c["segment"]], conf(c), first_url(c), DATE, "research/02_pharmacy_ruo_consumer.json"])
seen = {r[0].split(" (")[0].lower() for r in rows01}
for x in d3:
    nm = x["company"]
    if nm.startswith("None") or "not a company" in nm or nm.split(" (")[0].lower() in seen or any(nm.split(" (")[0].lower().split("/")[0].strip() in s for s in seen): continue
    seen.add(nm.split(" (")[0].lower())
    rows01.append([nm, x.get("company_url"), x["company_segment"], LEGAL.get(x["company_segment"], ""), x.get("nature_of_involvement"), "", "", x.get("celebrity") + " (" + x.get("relationship_class") + ")", x.get("audience_fit"), "", "", x.get("celebrity") + ": " + x.get("relationship_class"), x.get("marketing_impact_observed"), "", x.get("public_role"), "", "", "", x.get("audience_size"), "", "", x.get("reputation_or_concentration_risk"), "", "Celebrity-associated company added from 04_Celebrity_Founder_Database; not ranked (no product/pricing audit performed).", "Low-Medium", (x.get("evidence_urls") or [""])[0], DATE, "research/03_celebrity_database.json"])
with open(f"{O}/01_Peptide_Market_Landscape.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(cols)
    for r in rows01: w.writerow([S(v) for v in r])
print("01 rows", len(rows01))

# ---------------- 02 Website audit workbook
wb = Workbook()
readme(wb, "02 — Competitor Website, UI/UX and Conversion Audit (prepared 2026-09-02)", [
 LIMIT,
 "Sheets: Framework (23 criteria, 1-5 scale, definitions) · Scores (matrix; blank = not assessed) · Score Notes (justification per criterion) · Company Profiles · Customer Journey & Medical Experience (clinician-led companies) · Messaging & Claims · Acquisition Channels · Segment Rankings · Sources.",
 "Scale: 1 = poor / absent; 2 = weak; 3 = adequate; 4 = strong; 5 = best-in-class. Scores present in this file were inferred from third-party descriptions and search excerpts and carry the confidence flag shown; treat them as hypotheses to confirm in a rendered-page audit (checklist in Framework sheet).",
 "Companion files: 01_Peptide_Market_Landscape.csv, 03_Peptide_Product_and_Pricing_Audit.xlsx, 08_ReserveClinic_Evidence_Ledger.xlsx (all source URLs).",
])
CRIT = [("homepage_clarity","Homepage clarity","Can a first-time visitor state what the company does and for whom within 5 seconds?"),("value_proposition","Value proposition","Is the core promise specific, credible and differentiated?"),("brand_identity","Brand identity","Coherent, premium, category-appropriate identity"),("visual_hierarchy","Visual hierarchy","Typography, spacing and color guide attention to key actions"),("mobile_responsiveness","Mobile responsiveness","Layouts, tap targets and flows work on phones"),("navigation_ia","Navigation & information architecture","Findability; logical grouping; search/filtering where needed"),("product_discovery","Product/program discovery","Ease of finding relevant program or product"),("educational_content","Educational content","Depth, accuracy, review and citations"),("trust_signals","Trust signals","Licensure, pharmacy disclosure, security, reviews"),("medical_authority","Medical authority","Named clinicians, credentials, governance"),("founder_clinician_presentation","Founder & clinician presentation","Quality and transparency of founder/clinician presentation"),("testimonials_use","Testimonials & before/after","Compliant use; typicality disclosures; no B/A imagery"),("pricing_transparency","Pricing transparency","All-in price visible before intake; fees disclosed"),("faq_objection_handling","FAQ & objection handling","Anticipates safety, cost, eligibility, cancellation"),("ctas_lead_capture","CTAs & lead capture","Clear primary action; consented capture"),("intake_quiz","Intake / eligibility quiz","Safety screening quality and friction"),("checkout_subscription_flow","Checkout & subscription","Disclosure, consent, no dark patterns"),("account_refill_experience","Account & refill","Self-service, monitoring, refill criteria"),("support_visibility","Customer support visibility","Hours, channels, SLAs"),("shipping_returns_info","Shipping & returns","Clarity incl. Rx non-returnability, cold chain"),("accessibility","Accessibility","WCAG conformance signals"),("seo_fundamentals","SEO fundamentals","Structure, E-E-A-T, indexation"),("overall","Overall","Holistic conversion maturity")]
sheet(wb, "Framework", ["Key", "Criterion", "Definition", "Rendered-page checklist (to complete when access is available)"], [[k, n, d, "Load homepage, program/PDP, pricing, FAQ, quiz, checkout, account; capture screenshot with URL + date; score 1-5; record evidence."] for k, n, d in CRIT], widths=[26, 32, 60, 60])
allc = d1 + d2
hdr = ["Company", "Segment", "Audit confidence / status"] + [n for _, n, _ in CRIT]
rows = []; notes = []
for c in allc:
    wa = c.get("website_audit", {}) or {}
    status = S(wa.get("_status") or wa.get("audit_confidence") or ("Not assessed" if all(S((wa.get(k) or {}).get("score") if isinstance(wa.get(k), dict) else "") in ("", "unknown") for k, _, _ in CRIT) else "Low-confidence inference"))
    basis = S(wa.get("audit_basis", ""))
    row = [c["company"], c["segment"], (status[:120] + ("; " + basis[:120] if basis else ""))]
    for k, n, _ in CRIT:
        v = wa.get(k)
        sc = v.get("score") if isinstance(v, dict) else v
        try: sc = float(sc)
        except Exception: sc = ""
        row.append(sc)
        if isinstance(v, dict):
            notes.append([c["company"], c["segment"], n, sc, S(v.get("note") or v.get("justification"))])
    rows.append(row)
sheet(wb, "Scores", hdr, rows, widths=[34, 22, 40] + [12] * len(CRIT), note="Blank = not assessed (site could not be loaded). Numeric scores are low-confidence inferences from third-party descriptions; confirm with rendered-page audit. Scale 1-5.")
sheet(wb, "Score Notes", ["Company", "Segment", "Criterion", "Score", "Justification / note"], notes, widths=[34, 22, 30, 8, 90])
prof = []
for c in allc:
    prof.append([c["company"], c.get("url"), c["segment"], c.get("founded"), c.get("hq"), c.get("founders_leadership"), c.get("business_model"), c.get("target_market"), c.get("geography_availability"), c.get("clinical_pharmacy_partners") or c.get("partners"), c.get("celebrity_involvement"), c.get("funding_or_disclosed_performance"), c.get("customer_promise"), c.get("positioning"), c.get("traffic_proxy"), c.get("social_audience"), c.get("review_signal"), c.get("advertising_visibility"), c.get("regulatory_notes", ""), first_url(c), DATE])
sheet(wb, "Company Profiles", ["Company", "URL", "Segment", "Founded", "HQ", "Founders / leadership", "Business model", "Target market", "Geographic availability", "Clinical / pharmacy / fulfillment partners", "Celebrity involvement", "Funding or disclosed performance", "Primary customer promise", "Brand positioning", "Traffic proxy", "Social audience", "Review signal", "Advertising visibility", "Regulatory notes", "Primary evidence URL", "Access date"], prof, widths=[30, 28, 22, 10, 18, 30, 40, 28, 22, 30, 28, 34, 30, 30, 22, 22, 22, 22, 40, 34, 11])
cj = [[c["company"], c.get("customer_journey_notes"), c.get("medical_experience_notes")] for c in d1]
sheet(wb, "Journey & Medical", ["Company", "Customer journey notes (education → discovery → eligibility → intake → consult → Rx → payment → fulfillment → refill → cancel; friction, dark patterns, disclosures)", "Medical-experience notes (eligibility determination, clinician review, labs, contraindications, consent, AE instructions, monitoring, refill criteria, credentials, state availability)"], cj, widths=[30, 80, 80], note="Clinician-led companies only. Notes derive from third-party descriptions and search excerpts; no accounts were created and no health information was submitted.")
mc = []
for c in allc:
    m = c.get("messaging_claims")
    if isinstance(m, dict): mc.append([c["company"], c["segment"], m.get("headline"), m.get("core_promises"), m.get("risky_claims"), m.get("disclosures")])
    else: mc.append([c["company"], c["segment"], "", S(m), S(c.get("risk_notes")), ""])
sheet(wb, "Messaging & Claims", ["Company", "Segment", "Headline (paraphrase)", "Core promises / benefit language", "Risky, unsupported or FDA/FTC-sensitive claims", "Disclosures observed"], mc, widths=[30, 22, 30, 50, 50, 30])
sheet(wb, "Acquisition Channels", ["Company", "Segment", "Channels observed"], [[c["company"], c["segment"], S(c.get("acquisition_channels_observed"))] for c in allc], widths=[30, 22, 90])
rk = []
for seg, lst in RANKS.items():
    for i, k in enumerate(lst, 1):
        full = next((c["company"] for c in allc if k.lower() in c["company"].lower() and c["segment"] == seg), k)
        rk.append([seg, i, full, RATIONALE[seg]])
sheet(wb, "Segment Rankings", ["Segment", "Rank within segment", "Company", "Method"], rk, widths=[30, 10, 40, 90])
src = []
for c in allc:
    for s in c.get("sources", []) or []:
        if isinstance(s, dict): src.append([c["company"], s.get("claim"), s.get("url"), s.get("access_date", DATE), s.get("label")])
sheet(wb, "Sources", ["Company", "Claim", "URL", "Access date", "Label"], src, widths=[30, 70, 60, 11, 22])
wb.save(f"{O}/02_Competitor_Website_Audit.xlsx"); print("02 saved", len(rows), "companies", len(src), "sources")

# ---------------- 03 Product & pricing
wb = Workbook()
readme(wb, "03 — Peptide Product, Program, Pricing and Marketing Audit (prepared 2026-09-02)", [
 LIMIT,
 "Sheets: Product Inventory (every material offering found, with regulatory classification) · Price Bands (observed bands by ingredient/program; third-party unless noted) · Monetization Patterns · Claims Classification · Regulatory Legend · Pricing Research Plan · Sources.",
 "Dosage protocols are intentionally not reproduced. Regulatory classifications are analyst assessments against FDA 503A bulks-list status as of 2026-09-02 and require counsel confirmation.",
])
PK = ["product_or_program", "active_ingredient", "product_category", "claimed_purpose", "delivery_method", "dosage_form", "rx_required", "consult_required", "displayed_price", "consult_fee", "membership_fee", "subscription_terms", "est_monthly_recurring_cost", "bundles", "sourcing_disclosure", "coa_or_testing_claims", "pharmacy_disclosure", "availability_limits", "promotional_claims", "safety_disclosures", "regulatory_classification", "evidence_url", "access_date"]
inv = []
for c in allc:
    for p in c.get("products", []) or []:
        inv.append([c["company"], c["segment"]] + [S(p.get(k, "")) for k in PK])
sheet(wb, "Product Inventory", ["Company", "Segment", "Product / program", "Active ingredient", "Category", "Claimed purpose", "Delivery method", "Dosage form", "Rx required", "Consult required", "Displayed price", "Consult fee", "Membership fee", "Subscription terms", "Est. monthly recurring cost", "Bundles / protocols", "Sourcing disclosure", "COA / testing claims", "Pharmacy / manufacturer disclosure", "Availability limits", "Promotional claims (paraphrase)", "Safety disclosures", "Regulatory classification", "Evidence URL", "Access date"], inv, widths=[28, 20, 30, 26, 20, 30, 16, 14, 12, 14, 22, 14, 14, 22, 18, 20, 24, 24, 24, 22, 34, 24, 30, 36, 11])
bands = [
 ["Sermorelin (compounded, telehealth)", "Rx / compounded", "$79–$400 per month; telehealth cluster $126–$225 (Eden ~$126 first month, Ivím $129 ODT, Hone $130 + $149 membership, Maximus <$175, Joi/Blokes ~$199)", "Third-party estimate / company-reported", "research/01 (Eden, Ivím, Hone, Maximus, Joi)"],
 ["NAD+ injections (compounded)", "Rx / compounded", "$119–$395 per month (Ivím $249; Eden $15/shot); benchmarks $150–$400", "Third-party estimate", "research/01, research/06"],
 ["Tesamorelin (Egrifta; FDA-approved, off-label longevity use)", "Rx / approved drug", "$300–$700 per month (benchmark)", "Third-party estimate (unverified)", "research/06"],
 ["CJC-1295 / ipamorelin (compounded; ipamorelin Category 1)", "Rx / compounded", "$200–$450 per month (benchmark)", "Third-party estimate (unverified)", "research/06"],
 ["BPC-157 (transitional FDA status; Category 2 removal Apr 2026, PCAC vote Jul 2026, final rule pending)", "Unclear / transitional", "~$350 per 15 mg vial (Marek); bundled $150–$400/mo (Peter MD); Defy from $160/mo across 34 peptides", "Third-party estimate", "research/01"],
 ["Compounded semaglutide (wind-down)", "Rx / compounded — contracting", "$149–$399 per month (Henry Meds $149 oral; Eden $196–$296)", "Third-party estimate", "research/01, research/06"],
 ["Compounded tirzepatide (wind-down)", "Rx / compounded — contracting", "$299–$599 per month", "Third-party estimate (unverified)", "research/06"],
 ["Branded GLP-1 direct-to-consumer", "Rx / approved drug", "LillyDirect $299–$449; NovoCare $199 intro → $349; Wegovy pill $149–$199; TrumpRx ~$350", "Third-party report", "research/05"],
 ["TRT programs (telehealth)", "Rx", "$79–$99 (Peter MD, TRT Nation) to $199 (Fountain) and $225–$350 (Marek)", "Third-party estimate", "research/01"],
 ["Telehealth memberships / care plans", "Service fee", "$39 → $99 (Eden), $74.99 (Ivím), $149 (Hone), $129–$199 (Lifeforce)", "Company-reported / third-party", "research/01"],
 ["Consultation fees", "Service fee", "$0 (bundled) to $150–$250 (Defy $250 initial / $150 follow-up)", "Third-party estimate", "research/01"],
 ["Diagnostics memberships", "Service fee", "Function Health $499/yr; Superpower $199–$499/yr; InsideTracker $149/yr (+$340 Ultimate); SiPhox from $99–$125", "Company-reported / third-party", "research/02"],
 ["Longevity clinic memberships", "Service fee", "Fountain Life $10.5K–$85K per year", "Third-party report", "research/03"],
 ["Consumer cosmetic peptides", "Cosmetic", "The Ordinary $30.90–$32; Medik8 Liquid Peptides $66; Drunk Elephant Protini ~$68", "Company-reported / third-party", "research/02"],
 ["Collagen peptides (supplement)", "Supplement", "Vital Proteins promotional $16–$20; Momentous premium tier", "Third-party", "research/02"],
 ["Peptide clinician education", "Education", "SSRP Pro $999.50/yr or $99.95/mo; Foundations $99–$299; A4M ~$500/module (partner pricing)", "Company-reported / third-party", "research/02"],
 ["RUO peptide vendor prices", "RUO (not a lawful consumer product)", "NOT RETRIEVED (sites blocked); no prices recorded to avoid fabrication", "n/a", "research/02"],
]
sheet(wb, "Price Bands", ["Offering", "Classification", "Observed price band (2026)", "Evidence label", "Source file"], bands, widths=[44, 26, 80, 30, 26])
mon = [
 ["One-time purchases", "Consumer supplements/cosmetics; RUO vendors (not lawful for human use); per-vial peptide pricing at premium clinics (Marek)", "Simple; low LTV unless repeat", "Non-Rx shop only"],
 ["Memberships / monthly care plans", "Eden $39→$99; Ivím $74.99; Hone $149; Lifeforce $129–$199; Function $499/yr; Superpower $199–$499/yr", "Doubles effective cost for members; frequently criticized when hidden until checkout", "Adopt: one transparent membership tier with clear inclusions; never required to obtain Rx pricing"],
 ["Consultation fees", "$0 bundled (Hims-style) to $250 (Defy)", "Upfront fee reduces low-intent intakes but raises abandonment", "Adopt: modest fee, refundable if not eligible"],
 ["Lab fees", "Bundled (Hims planned) or separate $99–$600", "Lab-gating improves safety and retention", "Adopt: transparent lab pricing; required where clinically indicated"],
 ["Product subscriptions / auto-refill", "Standard across telehealth; ROSCA enforcement (NextMed order Dec 2025; FTC v. Hims 2026)", "Retention engine and top FTC exposure", "Adopt with pre-checkout disclosure, affirmative consent, 2-click cancel, renewal reminders"],
 ["Bundled programs / protocols", "'Stacks' at RUO vendors and some clinics", "Bundles of unapproved ingredients raise claim and safety risk", "Avoid ingredient 'stacks' marketing; program bundles = consult + monitoring + Rx if appropriate"],
 ["Financing", "BNPL (Affirm/Klarna) on supplements; rare on Rx", "Card-network and state rules on medical financing", "Defer; counsel review"],
 ["Discounts / introductory pricing", "Eden, Noom step-up pricing; Vital Proteins promos", "Introductory pricing with automatic step-up is a ROSCA disclosure hotspot", "Avoid step-up pricing; if used, disclose renewal price prominently"],
 ["Affiliates / influencer codes", "RUO vendors (Swiss Chems cited by FDA for social posts); consumer brands 10–30% commissions", "Affiliate marketing of Rx/RUO products draws FDA/FTC scrutiny", "Affiliates only for non-Rx products and education; disclosed; never tied to prescriptions"],
 ["Free-shipping thresholds", "Common in consumer brands", "Minor", "Adopt for non-Rx"],
 ["Cancellation terms", "Widely opaque; FTC actions cite hidden cancellation", "Trust and legal exposure", "Adopt: cancellation visible before purchase and in account"],
 ["Cross-sells / upsells", "Labs → Rx (Superpower, Hone); supplements alongside Rx (Hims)", "Effective when clinically justified", "Adopt education-led cross-sell; no upsell inside clinical consultation"],
 ["Customer lifetime-value logic", "Telehealth LTV driven by monthly retention (5–10% churn general; 8–15% GLP-1; 2–4% TRT per benchmarks)", "Retention > acquisition", "Model in 07_ReserveClinic_Financial_Model.xlsx"],
]
sheet(wb, "Monetization Patterns", ["Mechanism", "Observed in market", "Implication", "Reserve Clinic recommendation [SR]"], mon, widths=[30, 70, 50, 60])
claims = [
 ["Longevity / anti-aging", "NAD+, sermorelin, 'ageless' framing (Hims longevity, Lifeforce, Fountain Life, clinics)", "Off-label promotion of compounded drugs; FTC substantiation", "High", "Describe mechanism and evidence grade; no anti-aging outcome claims"],
 ["Weight-loss", "GLP-1 telehealth; Lemme 'GLP-1 Daily' supplement (class actions 2025)", "FDA warning letters Sept 2025–Jun 2026 cite 'same as Ozempic', FDA-approval implication, personalization claims", "High", "Route GLP-1 demand to branded DTC or clinician; no compounded GLP-1 marketing"],
 ["Recovery / healing", "BPC-157, TB-500 marketing at clinics and RUO vendors", "Unapproved ingredients; animal-data extrapolation", "High", "Library entry with honest evidence; not offered until lawful and evidence-supported"],
 ["Sexual health", "PT-141 (bremelanotide is FDA-approved as Vyleesi for premenopausal women with HSDD)", "Off-label male use; 'essentially a copy' compounding limits", "Medium", "Only via clinician within approved-drug framework; counsel review"],
 ["Performance / body composition", "TRT clinics; 'optimization' framing (Marek, Peter MD)", "Anti-doping, cardiovascular safety; testimonials of physique", "Medium", "No physique before/after; clinician-defined eligibility"],
 ["Biohacking / exclusivity", "'Guided optimization', 'peptide bros' culture, influencer stacks", "Reputation; attracts scrutiny", "Medium", "Avoid jargon; premium calm tone"],
 ["Clinician-authority messaging", "'Doctor-designed', named MDs (Defy, Hone) vs anonymous 'our doctors'", "Unsupported authority claims", "Low-Medium", "Name clinicians, licenses, review dates"],
 ["Social proof / testimonials", "Trustpilot widgets (Peter MD 4.8; TRT Nation divergence 1.9 vs 4.9), fabricated reviews cited in FTC NextMed order", "FTC fake-review rule; typicality", "High", "Verified reviews only; no B/A; typicality disclosure"],
 ["Urgency / scarcity", "Intro pricing, limited slots", "Dark pattern", "Medium", "None"],
 ["Guarantees", "'Results guaranteed' at some clinics", "Deceptive", "High", "None; satisfaction-based refund on non-Rx only"],
 ["Risk disclosures", "Often minimal on landing pages; hidden in FAQs", "Missing disclosures cited in FDA/FTC letters", "High", "Prominent compounded-not-approved, side-effect and emergency disclosures"],
 ["Research-use-only positioning", "RUO vendors; FDA March 2026 letters hold disclaimers irrelevant to intended use", "Unapproved new drug / misbranding", "Very High", "Never sell, link, or describe as available"],
]
sheet(wb, "Claims Classification", ["Claim family", "Observed patterns", "Concern (neutral)", "Risk", "Reserve Clinic copy rule [SR]"], claims, widths=[28, 70, 60, 10, 60])
leg = [
 ["FDA-approved drug", "Semaglutide (Ozempic/Wegovy), tirzepatide (Mounjaro/Zepbound), tesamorelin (Egrifta), bremelanotide (Vyleesi), teriparatide; testosterone products", "Prescription only; off-label promotion prohibited for brand"],
 ["Compounded preparation (503A/503B, permissible)", "Sermorelin, ipamorelin (Category 1), NAD+, glutathione; approved-drug APIs subject to 'essentially a copy' limits", "Not FDA-approved for safety/efficacy; disclosures required; patient-specific prescriptions (503A)"],
 ["FDA-restricted / transitional bulk substance", "BPC-157, KPV, TB-500, MOTS-c, semax, epitalon (Category 2 removal Apr 2026; PCAC recommendation Jul 2026; final rule pending); CJC-1295, GHK-Cu inj., AOD-9604, Ta1, selank, dihexa, DSIP, kisspeptin (moved out of Category 2, status pending)", "Compounding authority unsettled; exclude until final rulemaking; monitor monthly"],
 ["Category 2 (remains restricted)", "Melanotan II, GHRP-2, GHRP-6, LL-37, PEG-MGF", "Do not offer"],
 ["Supplement (DSHEA)", "Collagen peptides, amino acids, urolithin A, botanicals", "Structure-function claims with substantiation; no disease claims"],
 ["Cosmetic", "Topical signal/copper peptides (The Ordinary, Medik8, Drunk Elephant, Olay)", "Appearance claims only"],
 ["Research-use-only (RUO)", "Any peptide sold 'for research' to consumers", "FDA treats human-use marketing as unapproved drug; excluded"],
 ["Unclear / potentially misleading", "Oral 'BPC-157' capsules; 'GLP-1' named botanicals", "Not lawful dietary ingredients / misleading naming; excluded"],
]
sheet(wb, "Regulatory Legend", ["Classification", "Examples found", "Handling"], leg, widths=[36, 90, 70])
plan = [
 ["1", "Competitive price bands", "This workbook (Price Bands)", "Floor/ceiling per program", "Complete (third-party bands; re-verify)"],
 ["2", "Customer interviews", "20–30 interviews across three segments", "Value drivers, objections, willingness to pay", "Phase 1"],
 ["3", "Van Westendorp price sensitivity", "n≥200 per program via panel", "Acceptable price range; optimal price point", "Phase 1"],
 ["4", "MaxDiff on program components", "n≥200", "Bundle design (clinician access, labs, monitoring, education)", "Phase 1"],
 ["5", "Landing-page price tests", "All-in vs itemized; two price points", "Waitlist conversion as proxy", "Phase 3"],
 ["6", "Waitlist conversion analysis", "Cohorts by source", "Demand elasticity", "Phase 3"],
 ["7", "Clinician & fulfillment cost floor", "Signed contracts", "Gross margin ≥55%", "Phase 0–1"],
 ["8", "Regulatory & support cost loading", "Consent, monitoring, AE handling, compliance staffing", "Included in contribution", "Phase 1"],
 ["9", "Perceived value & premium test", "Concept test vs Hone/Lifeforce/Superpower", "Premium justification", "Phase 1"],
 ["10", "Decision & governance", "GM sets price; medical director and compliance sign-off; quarterly review", "Documented rationale", "Phase 2"],
]
sheet(wb, "Pricing Research Plan", ["Step", "Method", "Design", "Output", "Timing"], plan, widths=[6, 34, 50, 44, 20])
sheet(wb, "Sources", ["Company", "Claim", "URL", "Access date", "Label"], src, widths=[30, 70, 60, 11, 22])
wb.save(f"{O}/03_Peptide_Product_and_Pricing_Audit.xlsx"); print("03 saved", len(inv), "products")

# ---------------- 04 Celebrity database
wb = Workbook()
readme(wb, "04 — Celebrity Founder, Investor and Ambassador Database (prepared 2026-09-02)", [
 "Every relationship is classified as exactly one of: 1 Verified founder or cofounder · 2 Verified equity owner or investor · 3 Executive or board member · 4 Medical advisor · 5 Brand ambassador or spokesperson · 6 Recurring paid influencer · 7 Affiliate · 8 One-time campaign participant · 9 Unverified or rumored association. Appearing in advertising does not make someone a cofounder.",
 LIMIT.replace("Website UI/UX scores are either 'not assessed' or low-confidence inferences and must be re-scored from rendered pages before external use. ", "") + " Follower counts are approximate mid-2025/2026 figures from search excerpts and must be re-verified before external use.",
 "Sheets: Database (59 records) · Classification Summary (COUNTIF formulas) · Lessons · Principal Watchlist (Scott Disick, Dr. Michael Azziz — proposed, unverified) · Sources. Narrative analysis: research/03_celebrity_database.md.",
])
K3 = ["company", "company_url", "company_segment", "celebrity", "relationship_class", "nature_of_involvement", "announcement_date", "evidence_urls", "ownership_verified", "relationship_current", "public_role", "audience_fit", "audience_size", "marketing_impact_observed", "reputation_or_concentration_risk", "ftc_disclosure_notes", "peptide_relevance", "label", "access_date"]
H3 = ["Company", "Company URL", "Segment", "Celebrity / public figure", "Relationship class", "Nature of involvement", "Announcement date", "Evidence URLs", "Ownership verified", "Relationship current", "Public role", "Audience fit", "Audience size (approx.)", "Marketing impact observed", "Reputation / concentration risk", "FTC disclosure notes", "Peptide relevance", "Evidence label", "Access date"]
ws = sheet(wb, "Database", H3, [[S(x.get(k)) for k in K3] for x in d3], widths=[26, 26, 22, 24, 30, 50, 14, 50, 12, 12, 26, 30, 26, 44, 44, 40, 12, 22, 11])
cs = wb.create_sheet("Classification Summary")
cs["A1"] = "Relationship class counts (formulas reference the Database sheet)"; cs["A1"].font = T
classes = ["1 - Verified founder or cofounder", "2 - Verified equity owner or investor", "3 - Executive or board member", "4 - Medical advisor", "5 - Brand ambassador or spokesperson", "6 - Recurring paid influencer", "7 - Affiliate", "8 - One-time campaign participant", "9 - Unverified or rumored association"]
cs.cell(row=3, column=1, value="Class").font = H; cs.cell(row=3, column=1).fill = HF; cs.cell(row=3, column=2, value="Count").font = H; cs.cell(row=3, column=2).fill = HF
n3 = len(d3) + 1
for i, cl in enumerate(classes, 4):
    cs.cell(row=i, column=1, value=cl).font = B
    cs.cell(row=i, column=2, value=f'=COUNTIF(Database!$E$2:$E${n3},"{cl}*")').font = B
cs.cell(row=4 + len(classes), column=1, value="Total").font = BOLD; cs.cell(row=4 + len(classes), column=2, value=f"=SUM(B4:B{3 + len(classes)})").font = BOLD
cs.cell(row=6 + len(classes), column=1, value="Peptide-direct records").font = B; cs.cell(row=6 + len(classes), column=2, value=f'=COUNTIF(Database!$Q$2:$Q${n3},"direct*")').font = B
cs.cell(row=7 + len(classes), column=1, value="Adjacent records").font = B; cs.cell(row=7 + len(classes), column=2, value=f'=COUNTIF(Database!$Q$2:$Q${n3},"adjacent*")').font = B
cs.column_dimensions["A"].width = 44; cs.column_dimensions["B"].width = 12
lessons = [
 ["Measurable advantage", "Founder-owned, product-led, peptide-positioned (Rhode: peptide lip SKU launch hero; $1B e.l.f. sale May 2025; $20M media impact value in 48h); investor-as-distribution (AG1 $1.2B valuation with podcast-host cap table; Onnit/Rogan → Unilever 2021); cap-table halo (Function Health: Damon, Efron, Hart, Pascal; $2.5B Nov 2025); trust transfer (Prenuvo/Kim Kardashian unpaid post cited in $120M raise narrative); public-market proof (Oprah/WW: stock roughly doubled on 2015 entry, fell ~25% on Feb 2024 exit).", "Celebrity works when the celebrity owns equity, engages repeatedly in trusted long-form formats, and the product has a hero SKU aligned with a term the audience already searches."],
 ["Cosmetic involvement", "Ladder (LeBron James et al.): ~$4M revenue after two years despite >180M combined followers; SKKN by Kim folded June 2025; Kin Euphorics cofounder title granted years after founding; investor cohorts (Superpower, Neko, Tally) with no continuing marketing activity.", "Audience size without founder engagement or product differentiation does not convert."],
 ["Founder-market fit", "Strong: Halle Berry (menopause), Mark Hyman (practitioner), Peter Attia (physician), Jay Campbell/Hunter Williams (peptide audience). Weak/borrowed: 'GLP-1' naming on a botanical (Lemme); tech fame without health credibility (Neko/Ek).", "Test: does the person already talk about the compound class unpaid? Scott Disick's Feb 2025 Mounjaro disclosure is an organic GLP-1 signal, not a peptide-therapy signal."],
 ["Trust transfer & conversion", "Medical-credibility transfer converts fastest and breaks hardest (Prenuvo/Kardashian; 10X/Brecka drew physician and McGill OSS rebuttals). Repeated podcast recommendations convert (AG1, Onnit, Momentous); a single Super Bowl spot converted to traffic (Hims +650%) not durable trust.", "Pair any celebrity with named clinicians and standards content; use repeated long-form formats over one-off spectacles."],
 ["Reputation & regulatory risk", "Goop 2018 $145K settlement and substantiation injunction; Kim Kardashian SEC $1.26M (2022) for undisclosed payment despite #AD; Lemme GLP-1 Daily class actions (CA and NY, 2025) and rebrand to 'Lemme Reset'; Liver King deceptive-advertising suit; Hims 2025 senator letters; Dr. Oz/iHerb disclosure letter; deepfake celebrity GLP-1 ads.", "Disclosure of ownership and compensation in every endorsement; competent and reliable scientific evidence for every claim; no drug-like claims for non-drugs."],
 ["Key-person risk & brand dependence", "TB12 rebranded to TBRx after Brady–Guerrero split; 10X Health litigation after Brecka exit ($100M vs $13M suits, settled 2025); Blueprint inseparable from Bryan Johnson; WW priced Oprah's exit at ~25% of market cap; Bulletproof managed decoupling.", "Contract for key-person events (earnouts, morals and PED-disclosure clauses, testimonial ownership); build brand equity in clinicians and standards, not only the celebrity."],
 ["Regulatory scrutiny specific to peptides", "HHS/FDA 2026 reclassification (Feb 27 announcement; Apr 15 503A update; Jul 23–24 PCAC) is explicitly linked in press to influencer demand (Rogan, Brecka, Means, Buhler); any celebrity peptide launch will be read as MAHA-adjacent; RUO vendors remain exposed.", "Decide product category (compounded Rx via physician vs cosmetic/collagen 'peptide' vs never RUO) before any celebrity announcement."],
 ["Scott Disick specifics (proposed, unverified)", "On FTC's Sept 2017 21-person influencer warning-letter list; 2016 Bootea caption incident; TINA.org database; confirmed Mounjaro use (Feb 2025); reported promotional GLP-1 post (brand/disclosure unverified — potential category exclusivity); documented addiction history; ~27M Instagram followers (search excerpt). No public link to Reserve Clinic or any peptide company found.", "Disclosure-perfect execution, morality/exclusivity terms, and a medical-substance-first brand are preconditions; his audience is a launch-awareness asset, not a trust asset."],
]
sheet(wb, "Lessons", ["Theme", "Evidence (see Database and Sources)", "Implication for Reserve Clinic [SR]"], lessons, widths=[30, 100, 70])
watch = [[e.get("topic"), e.get("claim"), e.get("url"), e.get("access_date"), e.get("label"), e.get("confidence")] for e in d4["evidence"]]
sheet(wb, "Principal Watchlist", ["Topic", "Claim", "URL", "Access date", "Label", "Confidence (0-1)"], watch, widths=[30, 90, 60, 11, 18, 12], note="Scott Disick and Dr. Michael Azziz are PROPOSED principals; no public confirmation of any relationship to Reserve Clinic was found. 'search snippet' = seen only in a search-result excerpt; 'blocked' = source unreachable from the research environment. Full narrative: research/04_principal_verification.md.")
src3 = []
for x in d3:
    for u in x.get("evidence_urls") or []: src3.append([x["company"], x["celebrity"], u, x.get("access_date", DATE), x.get("label")])
sheet(wb, "Sources", ["Company", "Celebrity", "URL", "Access date", "Label"], src3, widths=[30, 26, 70, 11, 24])
wb.save(f"{O}/04_Celebrity_Founder_Database.xlsx"); print("04 saved", len(d3))

# ---------------- 08 Evidence ledger
wb = Workbook()
readme(wb, "08 — Reserve Clinic Evidence Ledger (prepared 2026-09-02)", [
 "Purpose: single traceable register of every source used across the market audit, celebrity investigation, principal verification, regulatory brief, market sizing and financial-model assumptions. Each entry carries a label: Verified fact · Company-reported claim · Third-party estimate/report · Analyst inference · Search snippet · Recalled/unverified · Strategic recommendation · Financial assumption.",
 LIMIT,
 "Market-size and benchmark records from research/06 were produced WITHOUT live access (search budget exhausted before that workstream began) and are labeled 'recalled_*_unverified'; they were cross-checked against sibling research where possible. They are provided for structure and order-of-magnitude only and must be re-verified before appearing in any investor document. The business plan PDF presents them with that caveat.",
 "Sheets: Evidence Ledger · Principal Verification · Market Size · Benchmarks · Compliance Matrix · Model Assumption Traceability · Verification Backlog.",
 "Naming note: the concept was renamed from the AminoLord working name to Reserve Clinic (ReserveClinic.com) after the research phase. Research files and their evidence rows keep the earlier name; domain and trademark checks recorded for aminolord.com/AMINOLORD are historical and must be re-run for reserveclinic.com and RESERVE CLINIC (DNS on 2026-09-02 shows reserveclinic.com resolving to registrar parking addresses).",
])
led = []; n = 0
def add(ws_, topic, claim, url, date, label, conf_=""):
    global n; n += 1; led.append([f"E{n:04d}", ws_, topic, claim, url, date or "", label, conf_])
for c in d1:
    for s in c.get("sources", []) or []:
        if isinstance(s, dict): add("WS1/2 Telehealth audit", c["company"], s.get("claim"), s.get("url"), s.get("access_date", DATE), s.get("label"))
for c in d2:
    for s in c.get("sources", []) or []:
        if isinstance(s, dict): add("WS1/2/3 Pharmacy-RUO-Consumer audit", c["company"], s.get("claim"), s.get("url"), s.get("access_date", DATE), s.get("label"))
for x in d3:
    for u in x.get("evidence_urls") or []: add("WS4 Celebrity database", f"{x['company']} / {x['celebrity']}", x.get("nature_of_involvement"), u, x.get("access_date", DATE), x.get("label"))
for e in d4["evidence"]: add("WS5 Principal verification", e.get("topic"), e.get("claim"), e.get("url"), e.get("access_date"), e.get("label"), e.get("confidence"))
for i in d5:
    for u in i.get("key_sources") or []: add("WS8 Regulatory", i["issue"], i.get("why_it_matters")[:300], u, DATE, "Third-party report / Verified fact (see research/05_regulatory_brief.md sources table)")
for m in d6m: add("WS1 Market size", m.get("category"), f"{m.get('market_definition')} — {m.get('geography')} {m.get('base_year')}: {m.get('base_value_usd')} → {m.get('forecast_year')}: {m.get('forecast_value_usd')} (CAGR {m.get('cagr')})", m.get("source_url"), m.get("access_date"), m.get("label"), m.get("confidence"))
for b in d6b: add("WS10 Benchmarks", b.get("metric"), f"{b.get('value')} {b.get('unit')} — {b.get('notes')}", b.get("source_url"), b.get("access_date"), b.get("label"), b.get("confidence"))
sheet(wb, "Evidence Ledger", ["ID", "Workstream", "Company / topic", "Claim", "URL", "Access date", "Label", "Confidence"], led, widths=[8, 26, 30, 80, 60, 11, 30, 10])
sheet(wb, "Principal Verification", ["Topic", "Claim", "URL", "Access date", "Label", "Confidence (0-1)"], watch, widths=[30, 90, 60, 11, 18, 12], note="Proposed principals only. No public confirmation of participation in Reserve Clinic was found for either individual. Items marked 'blocked' must be re-run (USPTO, NPI Registry, state medical boards, WHOIS).")
K6 = ["category", "market_definition", "geography", "base_year", "base_value_usd", "forecast_year", "forecast_value_usd", "cagr", "source_name", "source_url", "access_date", "reported_or_inferred", "applicability_to_Reserve Clinic", "reserveclinic_model_relevance", "label", "confidence", "notes", "access_note"]
sheet(wb, "Market Size", ["Category", "Market definition", "Geography", "Base year", "Base value (USD)", "Forecast year", "Forecast value (USD)", "CAGR", "Source", "Source URL", "Access date", "Reported / inferred", "Applicability to Reserve Clinic", "Model relevance", "Label", "Confidence", "Notes", "Access note"], [[S(m.get(k)) for k in K6] for m in d6m], widths=[22, 40, 10, 8, 16, 8, 16, 10, 30, 40, 10, 22, 60, 14, 26, 10, 40, 50], note="All market-size records are RECALLED and UNVERIFIED (see README). Values are ranges expressed as e.g. '~45-50e9' = $45-50 billion.")
K6b = ["metric", "value", "unit", "company_or_source", "source_url", "access_date", "label", "confidence", "notes", "access_note"]
sheet(wb, "Benchmarks", ["Metric", "Value", "Unit", "Company / source", "Source URL", "Access date", "Label", "Confidence", "Notes", "Access note"], [[S(b.get(k)) for k in K6b] for b in d6b], widths=[40, 18, 10, 36, 40, 10, 30, 10, 60, 50], note="Benchmarks are RECALLED and UNVERIFIED unless cross-referenced ([xref]) to sibling research; verify before use.")
K5 = ["issue", "why_it_matters", "models_affected", "risk_level", "owner", "professional_review", "pre_launch_action", "ongoing_control", "key_sources"]
sheet(wb, "Compliance Matrix", ["Issue", "Why it matters", "Business models affected", "Preliminary risk", "Required owner", "Professional review", "Pre-launch action", "Ongoing control", "Key sources"], [[S(i.get(k)) for k in K5] for i in d5], widths=[36, 70, 24, 10, 26, 30, 60, 50, 50], note="Preliminary; not legal advice. Models: education_commerce (A), telehealth_mso (B), premium_membership (C). Full brief: research/05_regulatory_brief.md.")
trace = [
 ["Program price $249–$299/mo", "Price bands: sermorelin $79–$400; NAD+ $119–$395; memberships $39–$199; premium positioning", "03 Price Bands; research/01", "Third-party estimate"],
 ["Monthly Rx churn 10–13%", "DTC health 5–10%/mo; GLP-1 8–15%; TRT 2–4%", "research/06 benchmarks (recalled, unverified)", "Recalled/unverified"],
 ["Pharmacy cost $85–$110 per member-month", "Peptide pharmacy COGS $30–$120/month; compounded telehealth GM 70–80%", "research/06 benchmarks (recalled, unverified)", "Recalled/unverified"],
 ["Clinician cost per async consult $35–$45", "Async visit cost $20–$60 via provider networks", "research/06 benchmarks", "Recalled/unverified"],
 ["Payment fees 2.7–3.2%", "Standard 2.9% + $0.30; high-risk 3.5–5%", "research/06; research/05 (Stripe/Mastercard BRAM)", "Third-party report"],
 ["Lab price $99–$139; cost $75–$85", "Wholesale 50–100 biomarker panels $60–$150; DTC panels $99–$300", "research/06", "Recalled/unverified"],
 ["Membership price $19–$39/mo", "Function $499/yr; Superpower $199–$499/yr; Lifeforce $129/mo; InsideTracker $149/yr", "research/01, research/02", "Company-reported / third-party"],
 ["CAC per Rx member (output) $250–$600 base", "Telehealth CAC $100–$300 (gross adds); Hims marketing ~40–48% of revenue", "research/06", "Recalled/unverified"],
 ["Celebrity royalty 3%; retainer $25–40k/mo; equity 5–12%", "Licensing royalties 5–15% of net sales; ambassador equity 1–5% + cash for A-list; Rhode/Onnit/Oprah precedents", "research/06; research/03", "Third-party / analyst inference"],
 ["Celebrity traffic uplift 1.10–1.30x", "Hims Super Bowl +650% traffic (one-off); Rhode/Function launch demand; durability unproven", "research/03", "Third-party report / analyst inference"],
 ["Legal $40–60k/mo months 1–3", "MSO/PC setup $50–150k; multi-state CPOM memos", "research/06; research/05", "Recalled/unverified"],
 ["Insurance $6–10k/mo", "Entity telemed liability $10–40k/yr; product liability $2–15k/yr; cyber $2–10k/yr; MD malpractice $5–15k/yr", "research/06", "Recalled/unverified"],
 ["Technology $20–45k/mo post-launch", "Telehealth vendors $25–75/visit + $2–10k/mo minimums; Shopify Plus $2,300/mo; EHR per-provider fees", "research/06; 05_ReserveClinic_Website_Strategy.md", "Recalled/unverified"],
 ["State coverage 35–45% at launch", "CPOM/async restrictions by state; pharmacy licensure", "research/05", "Analyst inference"],
]
sheet(wb, "Model Assumption Traceability", ["Model assumption (07_ReserveClinic_Financial_Model.xlsx)", "Benchmark basis", "Evidence location", "Label"], trace, widths=[46, 80, 46, 26])
backlog = [
 ["P1", "Re-run all company website audits from rendered pages with screenshots (23 criteria × 75 companies)", "Egress blocked", "Product/UX"],
 ["P1", "Hims & Hers Q4-2025–Q2-2026 releases; FY2026 guidance; peptide launch SKUs/pricing", "Not accessed", "Finance"],
 ["P1", "FDA 503A bulks-list page and PCAC Jul 2026 materials; any final rule on BPC-157 et al.", "Snippet-based", "Regulatory counsel"],
 ["P1", "USPTO clearance for RESERVE CLINIC (Classes 5, 35, 44, 25) and counsel opinion on 'clinic' naming by entity/state; Disick/Talentless marks", "Blocked", "IP counsel"],
 ["P1", "NPI Registry, state medical boards, ABIM for the intended 'Dr. Michael Azziz'", "Blocked", "Founder / counsel"],
 ["P1", "Identify the Yahoo-reported Disick GLP-1 promotional post: brand, date, paid status, exclusivity", "Headline only", "Brand / counsel"],
 ["P1", "reserveclinic.com WHOIS/RDAP, registrar, Wayback; social handles", "Blocked", "Ops"],
 ["P2", "Traffic proxies (Similarweb/Semrush), Meta Ad Library, Google Ads Transparency, Trustpilot/BBB counts for top 30", "Not retrievable", "Growth"],
 ["P2", "Syndicated market reports (Grand View, Precedence, Fortune BI, McKinsey Future of Wellness 2025/26, KFF GLP-1 polls)", "Recalled only", "Strategy"],
 ["P2", "RUO vendor price grid (for risk monitoring only)", "Not retrieved", "Compliance"],
 ["P2", "Vendor pricing (Wheel, OpenLoop, SteadyMD, Healthie, LegitScript, Stripe, Shopify)", "Recalled only", "Product/Finance"],
 ["P3", "Google Trends series for 'peptides', 'BPC-157', 'sermorelin', 'peptide therapy near me', 'GLP-1'", "Not fetched", "Growth"],
 ["P3", "Kardashian-Jenner brand disclosure-practice audit (SKKN, Poosh, Khy)", "Not completed", "Brand/compliance"],
]
sheet(wb, "Verification Backlog", ["Priority", "Item", "Why outstanding", "Owner"], backlog, widths=[8, 100, 20, 20])
wb.save(f"{O}/08_ReserveClinic_Evidence_Ledger.xlsx"); print("08 saved", len(led), "ledger rows")
