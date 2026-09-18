"""Pull the recalculated model into build/margin_values.json for charting.
The workbook is the source of truth; nothing is recomputed in Python."""
import json, os, sys
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
XL = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "outputs", "09_ReserveClinic_Margin_and_Growth_Model.xlsx")
wb = load_workbook(XL, data_only=True)
lay = json.load(open(os.path.join(HERE, "margin_layout.json")))
R, AR = lay["line_rows"], lay["assumption_rows"]
SCN = ["Conservative", "Baseline", "Aggressive", "No_Partner"]
KEYS = ["sessions", "partner_sessions", "partner_impressions", "organic_sessions", "paid_budget",
        "search_available", "search_spend", "search_sessions", "retarget_pool", "retarget_available",
        "retarget_spend", "retarget_sessions", "prospect_spend", "prospect_sessions", "paid_spend",
        "paid_sessions", "eff_cps", "nonrx_demand", "nonrx_capacity", "nonrx_new_cust", "nonrx_lost",
        "nonrx_orders", "nonrx_subs_active", "nonrx_rev_net", "nonrx_gp", "rx_active", "rx_new",
        "rx_rev_net", "rx_gp", "rx_intakes", "rx_weighted_sessions", "net_rev", "gross_profit",
        "gross_margin", "total_opex", "ebitda", "ebitda_margin", "cum_cash", "state_cov",
        "rx_contrib_mm", "nonrx_contrib_order", "rx_ltv", "nonrx_ltv", "new_cust_all",
        "cac_blended", "cac_search", "cac_retarget", "cac_prospect", "cac_paid_nonrx",
        "payback_months", "ox_paid", "ox_creative", "ox_team", "ox_launch"]
out = {"monthly": {}, "summary": {}, "assum": {}, "sens": {}}
for scn in SCN:
    ws = wb[f"Model_{scn}"]
    out["monthly"][scn] = {k: [ws.cell(row=R[k], column=2 + m).value or 0 for m in range(1, 37)] for k in KEYS}
sm = wb["Summary"]
for r in range(5, 45):
    lab = sm.cell(row=r, column=1).value
    if not lab: continue
    out["summary"][lab] = {s: sm.cell(row=r, column=c).value for s, c in zip(SCN, (2, 3, 4, 5))}
aw = wb["Assumptions"]
for k, r in AR.items():
    out["assum"][k] = {"label": aw.cell(row=r, column=2).value, "unit": aw.cell(row=r, column=3).value,
                       "Conservative": aw.cell(row=r, column=4).value, "Baseline": aw.cell(row=r, column=5).value,
                       "Aggressive": aw.cell(row=r, column=6).value, "No_Partner": aw.cell(row=r, column=7).value,
                       "conf": aw.cell(row=r, column=9).value, "sens": aw.cell(row=r, column=10).value}
sv = wb["Sensitivity"]
out["sens"]["contrib_mm"] = sv["D5"].value
out["sens"]["ltvcac"] = {"churns": [sv.cell(row=i, column=1).value for i in range(9, 15)],
                         "cacs": [sv.cell(row=8, column=j).value for j in range(2, 8)],
                         "grid": [[sv.cell(row=i, column=j).value for j in range(2, 8)] for i in range(9, 15)]}
out["sens"]["reach"] = {"reaches": [sv.cell(row=i, column=1).value for i in range(22, 27)],
                        "ctrs": [sv.cell(row=21, column=j).value for j in range(2, 7)],
                        "grid": [[sv.cell(row=i, column=j).value for j in range(2, 7)] for i in range(22, 27)]}
out["sens"]["caccurve"] = [[sv.cell(row=r, column=j).value for j in range(1, 6)] for r in range(46, 54)]
json.dump(out, open(os.path.join(HERE, "margin_values.json"), "w"))
print(f"extracted {len(SCN)} scenarios, {len(KEYS)} lines, {len(out['summary'])} summary metrics")
