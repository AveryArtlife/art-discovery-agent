"""Build 07_AminoLord_Financial_Model.xlsx — driver-based, editable, 3 scenarios, 36 months.
All numbers on the Assumptions sheet are FINANCIAL ASSUMPTIONS (blue). Every model cell is a formula."""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

OUT = sys.argv[1] if len(sys.argv) > 1 else "outputs/07_AminoLord_Financial_Model.xlsx"
MONTHS = 36
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF", size=10)
BLACK = Font(name=FONT, color="000000", size=10)
GREEN = Font(name=FONT, color="008000", size=10)
BOLD = Font(name=FONT, bold=True, size=10)
H1 = Font(name=FONT, bold=True, size=14)
H2 = Font(name=FONT, bold=True, size=11, color="FFFFFF")
YELLOW = PatternFill("solid", fgColor="FFFF00")
HDR = PatternFill("solid", fgColor="1F3D33")
SUB = PatternFill("solid", fgColor="E8EFEA")
thin = Side(style="thin", color="BBBBBB")
CUR = '$#,##0;($#,##0);-'
PCT = '0.0%'
NUM = '#,##0;(#,##0);-'
DEC = '#,##0.00'

wb = Workbook()

# ------------------------------------------------------------------ Assumptions
# key: (label, unit, cons, base, up, rationale/source, confidence, sensitivity, validation)
A = [
 ("SECTION", "Timing"),
 ("launch_month", "Public launch month (model month #)", "month", 9, 8, 7, "Phase plan in 06_AminoLord_Launch_Plan.md: Phase 0-3 precede launch; conservative adds one month of slippage.", "Medium", "High", "Phase exit criteria"),
 ("beta_month", "Private beta start month", "month", 6, 5, 5, "Phase 2 begins after internal pilot.", "Medium", "Low", "Phase exit criteria"),
 ("SECTION", "Traffic"),
 ("organic_m1", "Organic + direct sessions in month 1", "sessions", 1500, 2500, 4000, "Prototype site with 20 library entries; comparable early-stage health content sites. Analyst assumption.", "Low", "Medium", "GA4 after month 1"),
 ("organic_growth", "Organic monthly growth rate", "%", 0.07, 0.09, 0.11, "Library-led SEO compounding; DTC health benchmarks show 8-15%/mo early growth from a small base (analyst inference from research/06 benchmarks).", "Low", "High", "Search Console trend, quarterly"),
 ("organic_cap", "Organic + direct sessions ceiling per month", "sessions", 150000, 300000, 400000, "Caps compounding growth; category-level informational search demand is finite (analyst inference).", "Low", "Medium", "Search Console"),
 ("cps", "Cost per paid session (blended)", "$", 2.30, 2.00, 1.80, "Health/wellness CPC and CPM proxies; Meta health-category restrictions raise cost. Analyst assumption.", "Medium", "High", "Ad platform reports"),
 ("paid_pre", "Paid media spend per month, months 1 to beta-1", "$", 10000, 20000, 30000, "Waitlist/education ads only.", "Medium", "Low", "Budget actuals"),
 ("paid_beta", "Paid media spend per month, beta to launch-1", "$", 30000, 60000, 100000, "Waitlist building (Phase 3).", "Medium", "Medium", "Budget actuals"),
 ("paid_launch", "Paid media spend, launch month", "$", 200000, 350000, 500000, "Phase 4 budget range.", "Medium", "Medium", "Budget actuals"),
 ("paid_post", "Paid media spend per month after launch (month launch+1)", "$", 100000, 180000, 250000, "Phase 5 operating range; scaled by growth below.", "Medium", "High", "Payback discipline gate"),
 ("paid_growth", "Monthly growth in post-launch paid spend", "%", 0.02, 0.03, 0.04, "Spend grows only while payback <= target.", "Medium", "Medium", "Monthly CAC review"),
 ("SECTION", "Funnel (prescription programs)"),
 ("quiz_start", "Quiz start rate (starts / sessions)", "%", 0.09, 0.10, 0.12, "Telehealth quiz-start benchmarks 5-10% (analyst inference).", "Medium", "High", "Analytics"),
 ("quiz_complete", "Quiz completion rate", "%", 0.60, 0.70, 0.75, "One-question-per-screen quizzes typically 55-75%.", "Medium", "Medium", "Analytics"),
 ("eligible_outcome", "Share of completes with an eligible outcome", "%", 0.50, 0.60, 0.65, "Red-flag and state screening remove 40-55%.", "Medium", "Medium", "Analytics"),
 ("state_cov_launch", "Share of US population in served states at launch", "%", 0.35, 0.40, 0.45, "4-6 pilot states incl. large states (TX, FL) if licensure achieved.", "Medium", "High", "State readiness tracker"),
 ("state_cov_step", "Monthly increase in population coverage after launch", "%", 0.02, 0.03, 0.035, "Tranche-based expansion.", "Medium", "Medium", "State readiness tracker"),
 ("state_cov_cap", "Maximum population coverage", "%", 0.75, 0.85, 0.90, "Some states excluded for CPOM/async restrictions.", "Medium", "Low", "Counsel memo"),
 ("intake_submit", "Intake submission rate (eligible → intake)", "%", 0.48, 0.55, 0.60, "ID verification and consent friction.", "Medium", "High", "Analytics"),
 ("consult_complete", "Consult completion rate", "%", 0.75, 0.82, 0.88, "Async review completes higher than video.", "Medium", "Medium", "EHR data"),
 ("rx_rate", "Clinician prescribing rate (clinical decision; not a target)", "%", 0.55, 0.62, 0.68, "Observed ranges for telehealth GLP-1/hormone programs 50-70% (third-party reports); monitored for safety only.", "Low", "High", "Medical director audit"),
 ("enroll_rate", "Enrollment (paid) rate after prescription", "%", 0.58, 0.65, 0.68, "Price shock at plan stage; all-in pricing shown earlier improves this.", "Medium", "High", "Analytics"),
 ("SECTION", "Pricing"),
 ("program_price", "All-in program price per member-month (blended)", "$", 249, 275, 299, "Competitive bands in 03_Peptide_Product_and_Pricing_Audit.xlsx: compounded peptide programs ~$150-$450/mo; premium positioning.", "Medium", "High", "Van Westendorp + tests"),
 ("consult_fee", "Consultation fee (charged at intake)", "$", 49, 79, 99, "Range $0-$150 among competitors; refundable if not eligible is common.", "Medium", "Medium", "Pricing test"),
 ("lab_attach", "Share of new intakes ordering labs", "%", 0.40, 0.50, 0.60, "Programs requiring labs (hormone/metabolic).", "Medium", "Medium", "EHR data"),
 ("lab_price", "Lab panel price charged", "$", 99, 119, 139, "Direct-pay panels via lab partners.", "Medium", "Low", "Vendor quotes"),
 ("member_price", "Membership price per month (education, navigation, member pricing)", "$", 19, 29, 39, "Function Health $499/yr, Superpower $199-499/yr and Lifeforce $129/mo bracket the range; AminoLord tier is lighter.", "Medium", "Medium", "Pricing test"),
 ("member_attach_nonelig", "Membership attach among non-eligible quiz completers", "%", 0.04, 0.06, 0.08, "Analyst assumption.", "Low", "Medium", "Analytics"),
 ("member_attach_traffic", "Membership sign-ups per session (organic)", "%", 0.0005, 0.001, 0.0015, "Analyst assumption.", "Low", "Low", "Analytics"),
 ("shop_conv", "Non-Rx shop conversion (orders / sessions)", "%", 0.006, 0.009, 0.012, "DTC supplement conversion 0.5-1.5%.", "Medium", "Medium", "Shopify analytics"),
 ("shop_aov", "Non-Rx average order value", "$", 65, 75, 85, "Premium supplement/cosmetic AOV.", "Medium", "Low", "Shopify analytics"),
 ("shop_repeat", "Monthly reorder rate of prior shop customers", "%", 0.12, 0.15, 0.18, "Subscription attach on consumables.", "Low", "Medium", "Shopify analytics"),
 ("SECTION", "Retention"),
 ("rx_churn", "Monthly churn, prescription programs (blended)", "%", 0.13, 0.11, 0.10, "DTC health subscription churn 8-15%/mo; GLP-1 programs higher. Third-party benchmarks in research/06.", "Low", "High", "Cohort analysis"),
 ("member_churn", "Monthly churn, membership", "%", 0.08, 0.06, 0.05, "Content memberships 5-8%.", "Low", "Medium", "Cohort analysis"),
 ("SECTION", "Cost of revenue"),
 ("pharm_cost", "Pharmacy/medication cost per Rx member-month", "$", 110, 95, 85, "Compounded peptide and GLP-1 wholesale-to-telehealth cost bands (third-party reports); volume tiers reduce cost.", "Low", "High", "Pharmacy contracts"),
 ("clin_consult_cost", "Clinician cost per completed consult", "$", 45, 40, 35, "Async visit cost via provider networks ~$25-60.", "Medium", "Medium", "Vendor contracts"),
 ("clin_refill_cost", "Clinician cost per active Rx member-month (monitoring/refill review)", "$", 10, 8, 7, "Monthly check-in review.", "Medium", "Medium", "Vendor contracts"),
 ("ship_cost", "Fulfillment + shipping per Rx shipment (cold chain where needed)", "$", 18, 15, 13, "Pharmacy carrier programs.", "Medium", "Low", "Pharmacy contracts"),
 ("lab_cost", "Lab panel cost", "$", 85, 80, 75, "Direct-pay lab wholesale.", "Medium", "Low", "Vendor quotes"),
 ("pay_fee", "Payment processing fee (% of revenue)", "%", 0.032, 0.029, 0.027, "Card fees 2.7-3.2% + fixed; high-risk processors cost more.", "High", "Low", "Processor statements"),
 ("refund_pct", "Refunds and chargebacks (% of revenue)", "%", 0.05, 0.035, 0.025, "Benchmark 2-5%.", "Medium", "Medium", "Finance actuals"),
 ("support_cost", "Support cost per active member-month (non-clinical)", "$", 6, 5, 4, "Tooling + staffing per member.", "Medium", "Low", "Support actuals"),
 ("shop_cogs_pct", "Non-Rx product COGS (% of revenue)", "%", 0.40, 0.35, 0.32, "Premium supplements 30-40%.", "Medium", "Low", "Supplier quotes"),
 ("shop_fulfil", "Non-Rx fulfillment per order", "$", 9, 8, 7, "3PL pick/pack/ship.", "Medium", "Low", "3PL quotes"),
 ("member_cogs", "Membership delivery cost per member-month", "$", 3, 2, 2, "Platform + content amortization.", "Medium", "Low", "Actuals"),
 ("SECTION", "Operating expenses (monthly, fully loaded)"),
 ("team_p0", "Team cost per month, months 1-3", "$", 45000, 60000, 80000, "GM, product lead, compliance (fractional), engineering contractors.", "Medium", "Medium", "Hiring plan"),
 ("team_p1", "Team cost per month, month 4 to launch-1", "$", 90000, 120000, 160000, "Adds content, support, ops.", "Medium", "Medium", "Hiring plan"),
 ("team_p2", "Team cost per month, launch to month 12", "$", 140000, 180000, 240000, "Growth, support scaling.", "Medium", "Medium", "Hiring plan"),
 ("team_p3", "Team cost per month, months 13-24", "$", 160000, 260000, 340000, "", "Low", "Medium", "Hiring plan"),
 ("team_p4", "Team cost per month, months 25-36", "$", 190000, 340000, 450000, "", "Low", "Medium", "Hiring plan"),
 ("med_lead", "Medical leadership (medical director + advisory) per month", "$", 15000, 20000, 25000, "Fractional medical director retainers $10-25k/mo (analyst inference); compensation never tied to Rx volume.", "Medium", "Low", "Agreement"),
 ("tech_pre", "Technology & software per month, pre-launch", "$", 8000, 12000, 18000, "Vendor sandboxes, CMS, analytics, CMP.", "Medium", "Low", "Vendor quotes"),
 ("tech_post", "Technology & software per month, post-launch", "$", 20000, 30000, 45000, "Telehealth/EHR platform fees, Shopify, Stripe, warehouse, support tools.", "Medium", "Medium", "Vendor quotes"),
 ("build_capex", "Site & portal build per month, months 1-4 (one-time)", "$", 40000, 60000, 80000, "Option C build $150-300k in 05_AminoLord_Website_Strategy.md.", "Medium", "Low", "SOW"),
 ("legal_p0", "Legal, regulatory & compliance per month, months 1-3", "$", 40000, 50000, 60000, "Structuring, CPOM memos, formulary review, trademark, contracts.", "Medium", "Low", "Counsel estimates"),
 ("legal_post", "Legal & compliance per month thereafter", "$", 15000, 20000, 25000, "Ongoing review, state expansion, claims review.", "Medium", "Low", "Counsel estimates"),
 ("insurance", "Insurance per month (from launch)", "$", 6000, 8000, 10000, "Professional, product, cyber, D&O.", "Medium", "Low", "Broker quotes"),
 ("creative_pre", "Creative & content production per month, pre-launch", "$", 15000, 25000, 40000, "Library, video series, photography.", "Medium", "Low", "Agency quotes"),
 ("creative_post", "Creative & content production per month, post-launch", "$", 25000, 40000, 60000, "", "Medium", "Low", "Agency quotes"),
 ("launch_pr", "Launch PR/events one-time (launch month)", "$", 60000, 90000, 120000, "Phase 4 budget.", "Medium", "Low", "Agency quotes"),
 ("influ_pre", "Influencer/affiliate per month, beta to launch-1", "$", 10000, 25000, 50000, "Disclosed creator seeding.", "Medium", "Medium", "Contracts"),
 ("influ_post", "Influencer/affiliate per month, post-launch", "$", 20000, 50000, 100000, "", "Medium", "Medium", "Contracts"),
 ("SECTION", "Principal (celebrity) economics — hypothetical; no terms offered or accepted"),
 ("named_partner", "Named-partner model active? (1 = Scenario A, 0 = Scenario B)", "flag", 1, 1, 1, "Toggle to compare Scenario A vs B. Set to 0 for the independent-brand model.", "n/a", "High", "Founder decision"),
 ("celeb_cash", "Brand partner cash retainer per month (from beta month)", "$", 25000, 30000, 40000, "Hypothetical mixed structure; see Celebrity Economics sheet for alternatives.", "Low", "High", "Negotiation"),
 ("celeb_royalty", "Brand partner royalty (% of net revenue)", "%", 0.03, 0.03, 0.03, "Hypothetical; licensing royalties in consumer brands commonly 3-10% (third-party reports).", "Low", "High", "Negotiation"),
 ("celeb_equity", "Brand partner equity (% fully diluted; non-cash, for reference)", "%", 0.05, 0.08, 0.12, "Hypothetical.", "Low", "n/a", "Negotiation"),
 ("celeb_lift", "Traffic uplift multiplier when named-partner model is active (post-launch organic/direct)", "x", 1.10, 1.25, 1.30, "Celebrity-founded brands show launch-window demand spikes (see 04_Celebrity_Founder_Database.xlsx lessons); durability uncertain.", "Low", "High", "Brand-lift study"),
 ("SECTION", "Working capital"),
 ("inv_months", "Non-Rx inventory on hand (months of COGS)", "months", 2.0, 1.5, 1.5, "Initial stock.", "Medium", "Low", "Ops"),
 ("opening_cash", "Opening cash (for cash-balance line; capital requirement computed separately)", "$", 0, 0, 0, "Set to committed capital once known.", "n/a", "n/a", "Cap table"),
]

ws = wb.active; ws.title = "Assumptions"
ws["A1"] = "AminoLord — Financial Model Assumptions (all values are FINANCIAL ASSUMPTIONS unless a source is cited)"; ws["A1"].font = H1
ws["A2"] = "Blue = editable input. Three scenario columns drive Model_Conservative, Model_Base, Model_Upside. Change blue cells only; every other sheet recalculates. Prepared 2026-09-02; proposed concept."; ws["A2"].font = Font(name=FONT, italic=True, size=9)
hdrs = ["Key", "Assumption", "Unit", "Conservative", "Base", "Upside", "Rationale / source", "Confidence", "Sensitivity", "Validation method"]
for i, h in enumerate(hdrs, 1):
    c = ws.cell(row=4, column=i, value=h); c.font = H2; c.fill = HDR
ROW = {}
r = 5
for item in A:
    if item[0] == "SECTION":
        c = ws.cell(row=r, column=1, value=item[1]); c.font = BOLD
        for col in range(1, 11): ws.cell(row=r, column=col).fill = SUB
        r += 1; continue
    key, label, unit, cons, base, up, rat, conf, sens, val = item
    ws.cell(row=r, column=1, value=key).font = Font(name=FONT, size=8, color="777777")
    ws.cell(row=r, column=2, value=label).font = BLACK
    ws.cell(row=r, column=3, value=unit).font = BLACK
    for col, v in zip((4, 5, 6), (cons, base, up)):
        c = ws.cell(row=r, column=col, value=v); c.font = BLUE
        c.number_format = PCT if unit == "%" else (CUR if unit == "$" else (DEC if unit == "x" else '0.0' if unit=="months" else '0'))
        if sens == "High": c.fill = YELLOW
    ws.cell(row=r, column=7, value=rat).font = Font(name=FONT, size=9)
    ws.cell(row=r, column=8, value=conf).font = BLACK
    ws.cell(row=r, column=9, value=sens).font = BLACK
    ws.cell(row=r, column=10, value=val).font = Font(name=FONT, size=9)
    ROW[key] = r
    r += 1
for col, w in zip("ABCDEFGHIJ", (18, 62, 9, 14, 14, 14, 70, 11, 11, 24)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "D5"
ws.cell(row=r+1, column=2, value="Legend: yellow fill = high-sensitivity assumption to validate first. Scenario columns are independent; edit any of them.").font = Font(name=FONT, italic=True, size=9)

SCEN_COL = {"Conservative": "D", "Base": "E", "Upside": "F"}
def a(key, scen):
    return f"Assumptions!${SCEN_COL[scen]}${ROW[key]}"

# ------------------------------------------------------------------ Model sheets
FIRST = 4  # first month column index (D)
def mc(m): return L(FIRST + m - 1)  # month m -> column letter
LAST = mc(MONTHS)

def build_model(scen):
    s = wb.create_sheet(f"Model_{scen}")
    s["A1"] = f"AminoLord monthly model — {scen} scenario (all cells are formulas driven by the Assumptions sheet)"; s["A1"].font = H1
    s["A2"] = "Units: $ unless noted. Month 1 = start of Phase 0. Prescription programs run through an affiliated medical practice and licensed pharmacies; revenue recognition and merchant-of-record structure subject to counsel."; s["A2"].font = Font(name=FONT, italic=True, size=9)
    s.cell(row=3, column=1, value="Line").font = H2; s.cell(row=3, column=1).fill = HDR
    s.cell(row=3, column=2, value="Unit").font = H2; s.cell(row=3, column=2).fill = HDR
    s.cell(row=3, column=3, value="36-mo total").font = H2; s.cell(row=3, column=3).fill = HDR
    for m in range(1, MONTHS + 1):
        c = s.cell(row=3, column=FIRST + m - 1, value=m); c.font = H2; c.fill = HDR; c.alignment = Alignment(horizontal="center")
    s.column_dimensions["A"].width = 52; s.column_dimensions["B"].width = 9; s.column_dimensions["C"].width = 14
    for m in range(1, MONTHS + 1): s.column_dimensions[mc(m)].width = 11
    s.freeze_panes = "D4"
    R = {}
    state = {"r": 4}
    def section(title):
        r = state["r"]; c = s.cell(row=r, column=1, value=title); c.font = BOLD
        for col in range(1, FIRST + MONTHS): s.cell(row=r, column=col).fill = SUB
        state["r"] += 1
    def line(key, label, unit, fn, fmt=CUR, total=True, bold=False):
        """fn(m, col, prev_col) -> formula string (without '=')"""
        r = state["r"]; R[key] = r
        s.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
        s.cell(row=r, column=2, value=unit).font = BLACK
        for m in range(1, MONTHS + 1):
            col = mc(m); prev = mc(m - 1) if m > 1 else None
            c = s.cell(row=r, column=FIRST + m - 1, value="=" + fn(m, col, prev))
            c.font = BOLD if bold else BLACK; c.number_format = fmt
        if total:
            c = s.cell(row=r, column=3, value=f"=SUM({mc(1)}{r}:{LAST}{r})"); c.number_format = fmt; c.font = BOLD if bold else BLACK
        state["r"] += 1
        return r
    def ref(key, col): return f"{col}{R[key]}"
    A_ = lambda k: a(k, scen)

    section("Timing & coverage")
    line("month", "Model month", "#", lambda m, c, p: f"{c}$3", fmt='0', total=False)
    line("is_pre", "Pre-beta phase flag", "flag", lambda m, c, p: f"IF({c}$3<{A_('beta_month')},1,0)", fmt='0', total=False)
    line("is_beta", "Beta/waitlist phase flag", "flag", lambda m, c, p: f"IF(AND({c}$3>={A_('beta_month')},{c}$3<{A_('launch_month')}),1,0)", fmt='0', total=False)
    line("is_launch", "Launch month flag", "flag", lambda m, c, p: f"IF({c}$3={A_('launch_month')},1,0)", fmt='0', total=False)
    line("is_post", "Post-launch flag (incl. launch month)", "flag", lambda m, c, p: f"IF({c}$3>={A_('launch_month')},1,0)", fmt='0', total=False)
    line("state_cov", "Population coverage of served states (Rx)", "%",
         lambda m, c, p: f"IF({c}$3<{A_('beta_month')},0,IF({c}$3<{A_('launch_month')},{A_('state_cov_launch')}*0.5,MIN({A_('state_cov_cap')},{A_('state_cov_launch')}+{A_('state_cov_step')}*({c}$3-{A_('launch_month')}))))", fmt=PCT, total=False)

    section("Traffic")
    line("paid_spend", "Paid media spend", "$",
         lambda m, c, p: f"IF({c}$3<{A_('beta_month')},{A_('paid_pre')},IF({c}$3<{A_('launch_month')},{A_('paid_beta')},IF({c}$3={A_('launch_month')},{A_('paid_launch')},{A_('paid_post')}*(1+{A_('paid_growth')})^({c}$3-{A_('launch_month')}-1))))")
    line("paid_sessions", "Paid sessions", "sessions", lambda m, c, p: f"{ref('paid_spend', c)}/{A_('cps')}", fmt=NUM)
    line("organic_sessions", "Organic + direct sessions", "sessions",
         lambda m, c, p: f"MIN({A_('organic_cap')},{A_('organic_m1')}*(1+{A_('organic_growth')})^({c}$3-1))*IF(AND({ref('is_post', c)}=1,{A_('named_partner')}=1),{A_('celeb_lift')},1)", fmt=NUM)
    line("sessions", "Total sessions", "sessions", lambda m, c, p: f"{ref('paid_sessions', c)}+{ref('organic_sessions', c)}", fmt=NUM, bold=True)

    section("Prescription-program funnel (clinical decisions are independent; rates are monitored, not targeted)")
    line("quiz_starts", "Eligibility quiz starts", "#", lambda m, c, p: f"{ref('sessions', c)}*{A_('quiz_start')}*IF({c}$3<{A_('beta_month')},0,1)", fmt=NUM)
    line("quiz_completes", "Quiz completions", "#", lambda m, c, p: f"{ref('quiz_starts', c)}*{A_('quiz_complete')}", fmt=NUM)
    line("eligible", "Eligible outcomes (screen + state)", "#", lambda m, c, p: f"{ref('quiz_completes', c)}*{A_('eligible_outcome')}*{ref('state_cov', c)}", fmt=NUM)
    line("intakes", "Intakes submitted (consult fee charged)", "#", lambda m, c, p: f"{ref('eligible', c)}*{A_('intake_submit')}", fmt=NUM)
    line("consults", "Consults completed", "#", lambda m, c, p: f"{ref('intakes', c)}*{A_('consult_complete')}", fmt=NUM)
    line("prescribed", "Prescriptions issued (clinician decision)", "#", lambda m, c, p: f"{ref('consults', c)}*{A_('rx_rate')}", fmt=NUM)
    line("new_rx", "New paying Rx program members", "#", lambda m, c, p: f"{ref('prescribed', c)}*{A_('enroll_rate')}", fmt=NUM, bold=True)
    line("active_rx", "Active Rx program members (end of month)", "#", lambda m, c, p: (f"{ref('new_rx', c)}" if p is None else f"{ref('active_rx', p)}*(1-{A_('rx_churn')})+{ref('new_rx', c)}"), fmt=NUM, total=False, bold=True)
    line("churn_rx", "Rx members churned", "#", lambda m, c, p: ("0" if p is None else f"{ref('active_rx', p)}*{A_('rx_churn')}"), fmt=NUM)
    line("lab_orders", "Lab panels ordered", "#", lambda m, c, p: f"{ref('intakes', c)}*{A_('lab_attach')}", fmt=NUM)

    section("Membership and non-prescription commerce")
    line("new_members", "New membership sign-ups", "#", lambda m, c, p: f"({ref('quiz_completes', c)}-{ref('eligible', c)})*{A_('member_attach_nonelig')}+{ref('sessions', c)}*{A_('member_attach_traffic')}*IF({c}$3<{A_('beta_month')},0,1)", fmt=NUM)
    line("active_members", "Active members (end of month)", "#", lambda m, c, p: (f"{ref('new_members', c)}" if p is None else f"{ref('active_members', p)}*(1-{A_('member_churn')})+{ref('new_members', c)}"), fmt=NUM, total=False, bold=True)
    line("churn_members", "Members churned", "#", lambda m, c, p: ("0" if p is None else f"{ref('active_members', p)}*{A_('member_churn')}"), fmt=NUM)
    line("shop_new_orders", "Non-Rx first orders", "#", lambda m, c, p: f"{ref('sessions', c)}*{A_('shop_conv')}*IF({c}$3<{A_('beta_month')},0,1)", fmt=NUM)
    line("shop_customers", "Cumulative non-Rx customers", "#", lambda m, c, p: (f"{ref('shop_new_orders', c)}" if p is None else f"{ref('shop_customers', p)}+{ref('shop_new_orders', c)}"), fmt=NUM, total=False)
    line("shop_repeat_orders", "Non-Rx repeat orders", "#", lambda m, c, p: ("0" if p is None else f"{ref('shop_customers', p)}*{A_('shop_repeat')}"), fmt=NUM)
    line("shop_orders", "Total non-Rx orders", "#", lambda m, c, p: f"{ref('shop_new_orders', c)}+{ref('shop_repeat_orders', c)}", fmt=NUM)

    section("Revenue")
    line("rev_program", "Rx program revenue (all-in monthly)", "$", lambda m, c, p: f"{ref('active_rx', c)}*{A_('program_price')}")
    line("rev_consult", "Consultation revenue", "$", lambda m, c, p: f"{ref('intakes', c)}*{A_('consult_fee')}")
    line("rev_lab", "Lab revenue", "$", lambda m, c, p: f"{ref('lab_orders', c)}*{A_('lab_price')}")
    line("rev_member", "Membership revenue", "$", lambda m, c, p: f"{ref('active_members', c)}*{A_('member_price')}")
    line("rev_shop", "Non-Rx commerce revenue", "$", lambda m, c, p: f"{ref('shop_orders', c)}*{A_('shop_aov')}")
    line("rev_gross", "Gross revenue", "$", lambda m, c, p: f"{ref('rev_program', c)}+{ref('rev_consult', c)}+{ref('rev_lab', c)}+{ref('rev_member', c)}+{ref('rev_shop', c)}", bold=True)
    line("refunds", "Refunds & chargebacks", "$", lambda m, c, p: f"-{ref('rev_gross', c)}*{A_('refund_pct')}")
    line("rev_net", "Net revenue", "$", lambda m, c, p: f"{ref('rev_gross', c)}+{ref('refunds', c)}", bold=True)

    section("Cost of revenue")
    line("c_pharm", "Pharmacy / medication", "$", lambda m, c, p: f"-{ref('active_rx', c)}*{A_('pharm_cost')}")
    line("c_clin", "Clinical (consults + monitoring)", "$", lambda m, c, p: f"-({ref('consults', c)}*{A_('clin_consult_cost')}+{ref('active_rx', c)}*{A_('clin_refill_cost')})")
    line("c_lab", "Lab costs", "$", lambda m, c, p: f"-{ref('lab_orders', c)}*{A_('lab_cost')}")
    line("c_ship", "Rx fulfillment & shipping", "$", lambda m, c, p: f"-{ref('active_rx', c)}*{A_('ship_cost')}")
    line("c_shop", "Non-Rx COGS + fulfillment", "$", lambda m, c, p: f"-({ref('rev_shop', c)}*{A_('shop_cogs_pct')}+{ref('shop_orders', c)}*{A_('shop_fulfil')})")
    line("c_member", "Membership delivery", "$", lambda m, c, p: f"-{ref('active_members', c)}*{A_('member_cogs')}")
    line("c_pay", "Payment processing", "$", lambda m, c, p: f"-{ref('rev_gross', c)}*{A_('pay_fee')}")
    line("c_support", "Customer support (variable)", "$", lambda m, c, p: f"-({ref('active_rx', c)}+{ref('active_members', c)})*{A_('support_cost')}")
    line("cogs", "Total cost of revenue", "$", lambda m, c, p: f"SUM({ref('c_pharm', c)}:{ref('c_support', c)})", bold=True)
    line("gross_profit", "Gross profit", "$", lambda m, c, p: f"{ref('rev_net', c)}+{ref('cogs', c)}", bold=True)
    line("gm_pct", "Gross margin %", "%", lambda m, c, p: f"IF({ref('rev_net', c)}=0,0,{ref('gross_profit', c)}/{ref('rev_net', c)})", fmt=PCT, total=False)

    section("Acquisition and contribution")
    line("x_paid", "Paid media", "$", lambda m, c, p: f"-{ref('paid_spend', c)}")
    line("x_influ", "Influencer & affiliate", "$", lambda m, c, p: f"-IF({c}$3<{A_('beta_month')},0,IF({c}$3<{A_('launch_month')},{A_('influ_pre')},{A_('influ_post')}))")
    line("x_launch", "Launch PR/events (one-time)", "$", lambda m, c, p: f"-{ref('is_launch', c)}*{A_('launch_pr')}")
    line("x_celeb_cash", "Brand partner cash retainer (Scenario A only)", "$", lambda m, c, p: f"-{A_('named_partner')}*IF({c}$3<{A_('beta_month')},0,{A_('celeb_cash')})")
    line("x_celeb_roy", "Brand partner royalty (Scenario A only)", "$", lambda m, c, p: f"-{A_('named_partner')}*{ref('rev_net', c)}*{A_('celeb_royalty')}")
    line("acq_total", "Total acquisition & partner cost", "$", lambda m, c, p: f"SUM({ref('x_paid', c)}:{ref('x_celeb_roy', c)})", bold=True)
    line("contribution", "Contribution margin", "$", lambda m, c, p: f"{ref('gross_profit', c)}+{ref('acq_total', c)}", bold=True)
    line("new_paying", "New paying customers (Rx + members + first shop orders)", "#", lambda m, c, p: f"{ref('new_rx', c)}+{ref('new_members', c)}+{ref('shop_new_orders', c)}", fmt=NUM)
    line("cac", "Blended CAC (acquisition cost / new paying customers)", "$", lambda m, c, p: f"IF({ref('new_paying', c)}=0,0,-({ref('x_paid', c)}+{ref('x_influ', c)}+{ref('x_launch', c)})/{ref('new_paying', c)})", total=False)
    line("cac_rx", "CAC per new Rx member (all acquisition cost allocated to Rx)", "$", lambda m, c, p: f"IF({ref('new_rx', c)}=0,0,-({ref('x_paid', c)}+{ref('x_influ', c)}+{ref('x_launch', c)})/{ref('new_rx', c)})", total=False)

    section("Operating expenses")
    line("o_team", "Team (fully loaded)", "$", lambda m, c, p: f"-IF({c}$3<=3,{A_('team_p0')},IF({c}$3<{A_('launch_month')},{A_('team_p1')},IF({c}$3<=12,{A_('team_p2')},IF({c}$3<=24,{A_('team_p3')},{A_('team_p4')}))))")
    line("o_med", "Medical leadership & governance", "$", lambda m, c, p: f"-{A_('med_lead')}")
    line("o_tech", "Technology & software", "$", lambda m, c, p: f"-IF({ref('is_post', c)}=1,{A_('tech_post')},{A_('tech_pre')})")
    line("o_build", "Site & portal build (one-time)", "$", lambda m, c, p: f"-IF({c}$3<=4,{A_('build_capex')},0)")
    line("o_legal", "Legal, regulatory & compliance", "$", lambda m, c, p: f"-IF({c}$3<=3,{A_('legal_p0')},{A_('legal_post')})")
    line("o_ins", "Insurance", "$", lambda m, c, p: f"-IF({ref('is_post', c)}=1,{A_('insurance')},0)")
    line("o_creative", "Creative & content production", "$", lambda m, c, p: f"-IF({ref('is_post', c)}=1,{A_('creative_post')},{A_('creative_pre')})")
    line("opex", "Total operating expenses", "$", lambda m, c, p: f"SUM({ref('o_team', c)}:{ref('o_creative', c)})", bold=True)

    section("Profitability and cash")
    line("ebitda", "EBITDA / operating profit (loss)", "$", lambda m, c, p: f"{ref('contribution', c)}+{ref('opex', c)}", bold=True)
    line("wc", "Working capital investment (non-Rx inventory)", "$", lambda m, c, p: (f"-{A_('inv_months')}*ABS({ref('c_shop', c)})" if p is None else f"-{A_('inv_months')}*(ABS({ref('c_shop', c)})-ABS({ref('c_shop', p)}))"))
    line("cash_flow", "Net cash flow", "$", lambda m, c, p: f"{ref('ebitda', c)}+{ref('wc', c)}", bold=True)
    line("cum_cash", "Cumulative cash flow", "$", lambda m, c, p: (f"{ref('cash_flow', c)}" if p is None else f"{ref('cum_cash', p)}+{ref('cash_flow', c)}"), total=False, bold=True)
    line("cash_bal", "Cash balance (opening cash + cumulative)", "$", lambda m, c, p: f"{A_('opening_cash')}+{ref('cum_cash', c)}", total=False)
    line("be_helper", "Break-even helper (month # if EBITDA > 0, else 999)", "#", lambda m, c, p: f"IF({ref('ebitda', c)}>0,{c}$3,999)", fmt='0', total=False)

    section("Unit economics (steady-state view, per active Rx member)")
    line("arpu", "Revenue per active Rx member-month", "$", lambda m, c, p: f"IF({ref('active_rx', c)}=0,0,({ref('rev_program', c)})/{ref('active_rx', c)})", total=False)
    line("cm_per", "Contribution per Rx member-month before acquisition", "$", lambda m, c, p: f"{A_('program_price')}*(1-{A_('refund_pct')}-{A_('pay_fee')})-{A_('pharm_cost')}-{A_('clin_refill_cost')}-{A_('ship_cost')}-{A_('support_cost')}", total=False)
    line("ltv", "LTV (contribution-based) = contribution / churn", "$", lambda m, c, p: f"{ref('cm_per', c)}/{A_('rx_churn')}", total=False)
    line("ltv_cac", "LTV : CAC (Rx)", "x", lambda m, c, p: f"IF({ref('cac_rx', c)}=0,0,{ref('ltv', c)}/{ref('cac_rx', c)})", fmt='0.0x', total=False)
    line("payback", "CAC payback (months)", "months", lambda m, c, p: f"IF({ref('cm_per', c)}<=0,0,{ref('cac_rx', c)}/{ref('cm_per', c)})", fmt='0.0', total=False)
    line("rev_per_member", "Net revenue per active customer (Rx + membership)", "$", lambda m, c, p: f"IF(({ref('active_rx', c)}+{ref('active_members', c)})=0,0,{ref('rev_net', c)}/({ref('active_rx', c)}+{ref('active_members', c)}))", total=False)
    return s, R

MODELS = {}
for scen in ("Conservative", "Base", "Upside"):
    MODELS[scen] = build_model(scen)

# ------------------------------------------------------------------ Summary
sm = wb.create_sheet("Scenario Summary", 1)
sm["A1"] = "AminoLord — Scenario summary (all cells link to model sheets; green = cross-sheet link)"; sm["A1"].font = H1
sm["A2"] = "Proposed concept. Outputs are projections from labeled assumptions, not forecasts of results. Break-even = first month with positive EBITDA; capital required = most negative cumulative cash flow plus 15% contingency."; sm["A2"].font = Font(name=FONT, italic=True, size=9)
hdr = ["Metric", "Unit", "Conservative", "Base", "Upside"]
for i, h in enumerate(hdr, 1):
    c = sm.cell(row=4, column=i, value=h); c.font = H2; c.fill = HDR
def yr(sheet, row, y):
    c1 = mc(12 * (y - 1) + 1); c2 = mc(12 * y)
    return f"SUM({sheet}!{c1}{row}:{c2}{row})"
rows = []
def add(label, unit, fn, fmt=CUR):
    rows.append((label, unit, fn, fmt))
for y in (1, 2, 3):
    add(f"Year {y} net revenue", "$", lambda s, R, y=y: yr(f"Model_{s}", R["rev_net"], y))
for y in (1, 2, 3):
    add(f"Year {y} gross profit", "$", lambda s, R, y=y: yr(f"Model_{s}", R["gross_profit"], y))
for y in (1, 2, 3):
    add(f"Year {y} EBITDA", "$", lambda s, R, y=y: yr(f"Model_{s}", R["ebitda"], y))
add("Year 3 gross margin %", "%", lambda s, R: f"IF({yr(f'Model_{s}', R['rev_net'], 3)}=0,0,{yr(f'Model_{s}', R['gross_profit'], 3)}/{yr(f'Model_{s}', R['rev_net'], 3)})", PCT)
add("Active Rx members, month 12", "#", lambda s, R: f"Model_{s}!{mc(12)}{R['active_rx']}", NUM)
add("Active Rx members, month 24", "#", lambda s, R: f"Model_{s}!{mc(24)}{R['active_rx']}", NUM)
add("Active Rx members, month 36", "#", lambda s, R: f"Model_{s}!{mc(36)}{R['active_rx']}", NUM)
add("Active members (membership), month 36", "#", lambda s, R: f"Model_{s}!{mc(36)}{R['active_members']}", NUM)
add("Total sessions, 36 months", "#", lambda s, R: f"Model_{s}!C{R['sessions']}", NUM)
add("Blended CAC, month 12", "$", lambda s, R: f"Model_{s}!{mc(12)}{R['cac']}")
add("CAC per Rx member, month 12", "$", lambda s, R: f"Model_{s}!{mc(12)}{R['cac_rx']}")
add("CAC per Rx member, month 24", "$", lambda s, R: f"Model_{s}!{mc(24)}{R['cac_rx']}")
add("Contribution per Rx member-month", "$", lambda s, R: f"Model_{s}!{mc(12)}{R['cm_per']}")
add("LTV (contribution-based)", "$", lambda s, R: f"Model_{s}!{mc(12)}{R['ltv']}")
add("LTV : CAC, month 24", "x", lambda s, R: f"Model_{s}!{mc(24)}{R['ltv_cac']}", '0.0x')
add("CAC payback, month 24 (months)", "months", lambda s, R: f"Model_{s}!{mc(24)}{R['payback']}", '0.0')
add("Net revenue per active customer, month 36", "$", lambda s, R: f"Model_{s}!{mc(36)}{R['rev_per_member']}")
add("Break-even month (first positive EBITDA)", "month", lambda s, R: f"IF(MIN(Model_{s}!{mc(1)}{R['be_helper']}:{LAST}{R['be_helper']})=999,\"Not within 36 months\",MIN(Model_{s}!{mc(1)}{R['be_helper']}:{LAST}{R['be_helper']}))", '0')
add("Peak cumulative cash burn", "$", lambda s, R: f"MIN(Model_{s}!{mc(1)}{R['cum_cash']}:{LAST}{R['cum_cash']})")
add("Capital required (peak burn + 15% contingency)", "$", lambda s, R: f"-MIN(Model_{s}!{mc(1)}{R['cum_cash']}:{LAST}{R['cum_cash']})*1.15")
add("Cumulative cash flow, month 36", "$", lambda s, R: f"Model_{s}!{LAST}{R['cum_cash']}")
add("Brand partner cash + royalty paid, 36 months", "$", lambda s, R: f"-(Model_{s}!C{R['x_celeb_cash']}+Model_{s}!C{R['x_celeb_roy']})")
r = 5
for label, unit, fn, fmt in rows:
    sm.cell(row=r, column=1, value=label).font = BLACK
    sm.cell(row=r, column=2, value=unit).font = BLACK
    for col, scen in zip((3, 4, 5), ("Conservative", "Base", "Upside")):
        c = sm.cell(row=r, column=col, value="=" + fn(scen, MODELS[scen][1])); c.font = GREEN; c.number_format = fmt
    r += 1
for col, w in zip("ABCDE", (52, 9, 18, 18, 18)): sm.column_dimensions[col].width = w

# ------------------------------------------------------------------ Sensitivity
sv = wb.create_sheet("Sensitivity")
sv["A1"] = "Sensitivity analysis (Base scenario inputs; formulas recalculate when Assumptions change)"; sv["A1"].font = H1
sv["A2"] = "Table 1: LTV:CAC as a function of monthly Rx churn (rows) and CAC per Rx member (columns). Table 2: CAC payback months. Table 3: Year-3 Rx program revenue vs price and churn (steady-state approximation: month-36 new members / churn × price × 12)."; sv["A2"].font = Font(name=FONT, italic=True, size=9)
churns = [0.06, 0.08, 0.10, 0.11, 0.12, 0.14, 0.16]
cacs = [150, 200, 250, 300, 350, 400, 500]
B = "Base"
cm = f"(Assumptions!$E${ROW['program_price']}*(1-Assumptions!$E${ROW['refund_pct']}-Assumptions!$E${ROW['pay_fee']})-Assumptions!$E${ROW['pharm_cost']}-Assumptions!$E${ROW['clin_refill_cost']}-Assumptions!$E${ROW['ship_cost']}-Assumptions!$E${ROW['support_cost']})"
def grid(top, title, fn, fmt):
    sv.cell(row=top, column=1, value=title).font = BOLD
    sv.cell(row=top + 1, column=1, value="Churn ↓ / CAC →").font = BOLD
    for j, cac in enumerate(cacs):
        c = sv.cell(row=top + 1, column=2 + j, value=cac); c.font = BLUE; c.number_format = CUR
    for i, ch in enumerate(churns):
        c = sv.cell(row=top + 2 + i, column=1, value=ch); c.font = BLUE; c.number_format = PCT
        for j in range(len(cacs)):
            cell = sv.cell(row=top + 2 + i, column=2 + j, value="=" + fn(f"$A{top + 2 + i}", f"{L(2 + j)}${top + 1}"))
            cell.number_format = fmt; cell.font = BLACK
    return top + 2 + len(churns) + 2
nxt = grid(4, "Table 1 — LTV : CAC (contribution-based LTV = contribution per member-month / churn)", lambda ch, cac: f"({cm}/{ch})/{cac}", '0.0x')
nxt = grid(nxt, "Table 2 — CAC payback (months) = CAC / contribution per member-month", lambda ch, cac: f"{cac}/{cm}", '0.0')
# Table 3 price × churn
prices = [199, 225, 249, 275, 299, 325, 349]
top = nxt
sv.cell(row=top, column=1, value="Table 3 — Steady-state annual Rx program revenue ($) = (Base month-36 new Rx members / churn) × price × 12").font = BOLD
sv.cell(row=top + 1, column=1, value="Churn ↓ / Price →").font = BOLD
newrx = f"Model_Base!{LAST}{MODELS['Base'][1]['new_rx']}"
for j, pr in enumerate(prices):
    c = sv.cell(row=top + 1, column=2 + j, value=pr); c.font = BLUE; c.number_format = CUR
for i, ch in enumerate(churns):
    c = sv.cell(row=top + 2 + i, column=1, value=ch); c.font = BLUE; c.number_format = PCT
    for j in range(len(prices)):
        cell = sv.cell(row=top + 2 + i, column=2 + j, value=f"=({newrx}/$A{top + 2 + i})*{L(2 + j)}${top + 1}*12"); cell.number_format = CUR; cell.font = BLACK
top = top + 2 + len(churns) + 2
sv.cell(row=top, column=1, value="Table 4 — Single-variable sensitivity (Base): what each high-sensitivity input must be for LTV:CAC = 3.0 at the month-24 Base CAC").font = BOLD
cac24 = f"Model_Base!{mc(24)}{MODELS['Base'][1]['cac_rx']}"
sv.cell(row=top + 1, column=1, value="Month-24 Base CAC per Rx member").font = BLACK; c = sv.cell(row=top + 1, column=2, value=f"={cac24}"); c.font = GREEN; c.number_format = CUR
sv.cell(row=top + 2, column=1, value="Required contribution per member-month at Base churn for LTV:CAC = 3").font = BLACK; c = sv.cell(row=top + 2, column=2, value=f"=3*{cac24}*Assumptions!$E${ROW['rx_churn']}"); c.number_format = CUR
sv.cell(row=top + 3, column=1, value="Current Base contribution per member-month").font = BLACK; c = sv.cell(row=top + 3, column=2, value=f"={cm}"); c.number_format = CUR
sv.cell(row=top + 4, column=1, value="Maximum monthly churn for LTV:CAC = 3 at current contribution").font = BLACK; c = sv.cell(row=top + 4, column=2, value=f"=IF({cac24}=0,0,{cm}/(3*{cac24}))"); c.number_format = PCT
sv.cell(row=top + 5, column=1, value="Maximum CAC for LTV:CAC = 3 at current contribution and churn").font = BLACK; c = sv.cell(row=top + 5, column=2, value=f"={cm}/Assumptions!$E${ROW['rx_churn']}/3"); c.number_format = CUR
sv.column_dimensions["A"].width = 78
for j in range(2, 10): sv.column_dimensions[L(j)].width = 13

# ------------------------------------------------------------------ Celebrity economics
ce = wb.create_sheet("Celebrity Economics")
ce["A1"] = "Hypothetical brand-partner compensation structures (no terms have been offered, discussed, or accepted)"; ce["A1"].font = H1
ce["A2"] = "Uses Base-scenario net revenue. Equity value is illustrative and depends on an assumed exit multiple; it is not a valuation. All structures require FTC material-connection disclosure and counsel review of NIL, exclusivity, morality, and termination terms."; ce["A2"].font = Font(name=FONT, italic=True, size=9)
Rb = MODELS["Base"][1]
ce["A4"] = "Inputs"; ce["A4"].font = BOLD
inputs = [
 ("Base net revenue, year 1", f"={yr('Model_Base', Rb['rev_net'], 1)}", CUR, GREEN),
 ("Base net revenue, year 2", f"={yr('Model_Base', Rb['rev_net'], 2)}", CUR, GREEN),
 ("Base net revenue, year 3", f"={yr('Model_Base', Rb['rev_net'], 3)}", CUR, GREEN),
 ("Illustrative year-3 revenue multiple for equity value (assumption)", 3.0, '0.0x', BLUE),
 ("Structure 1 — Cash only: annual retainer", 600000, CUR, BLUE),
 ("Structure 2 — Equity only: fully diluted %", 0.10, PCT, BLUE),
 ("Structure 3 — Royalty only: % of net revenue", 0.06, PCT, BLUE),
 ("Structure 4 — Mixed: annual cash retainer", 360000, CUR, BLUE),
 ("Structure 4 — Mixed: royalty % of net revenue", 0.03, PCT, BLUE),
 ("Structure 4 — Mixed: equity % (fully diluted)", 0.05, PCT, BLUE),
 ("Structure 5 — Performance: royalty % above revenue hurdle", 0.08, PCT, BLUE),
 ("Structure 5 — Performance: annual net revenue hurdle", 5000000, CUR, BLUE),
]
for i, (lab, v, fmt, font) in enumerate(inputs):
    ce.cell(row=5 + i, column=1, value=lab).font = BLACK
    c = ce.cell(row=5 + i, column=2, value=v); c.number_format = fmt; c.font = font
IN = {lab: 5 + i for i, (lab, *_rest) in enumerate(inputs)}
top = 5 + len(inputs) + 2
ce.cell(row=top, column=1, value="Illustrative partner economics by structure").font = BOLD
for j, h in enumerate(["Structure", "Year 1 cash", "Year 2 cash", "Year 3 cash", "3-yr cash total", "Illustrative equity value at yr-3 multiple", "3-yr total value", "Cash cost as % of 3-yr revenue"]):
    c = ce.cell(row=top + 1, column=1 + j, value=h); c.font = H2; c.fill = HDR
rev = [f"$B${IN['Base net revenue, year 1']}", f"$B${IN['Base net revenue, year 2']}", f"$B${IN['Base net revenue, year 3']}"]
mult = f"$B${IN['Illustrative year-3 revenue multiple for equity value (assumption)']}"
structs = [
 ("1. Cash retainer only", [f"=$B${IN['Structure 1 — Cash only: annual retainer']}"] * 3, "0"),
 ("2. Equity only", ["=0"] * 3, f"=$B${IN['Structure 2 — Equity only: fully diluted %']}*{rev[2]}*{mult}"),
 ("3. Royalty only", [f"=$B${IN['Structure 3 — Royalty only: % of net revenue']}*{rv}" for rv in rev], "0"),
 ("4. Mixed (cash + royalty + equity)", [f"=$B${IN['Structure 4 — Mixed: annual cash retainer']}+$B${IN['Structure 4 — Mixed: royalty % of net revenue']}*{rv}" for rv in rev], f"=$B${IN['Structure 4 — Mixed: equity % (fully diluted)']}*{rev[2]}*{mult}"),
 ("5. Performance royalty above hurdle", [f"=$B${IN['Structure 5 — Performance: royalty % above revenue hurdle']}*MAX(0,{rv}-$B${IN['Structure 5 — Performance: annual net revenue hurdle']})" for rv in rev], "0"),
]
for i, (name, cash, eq) in enumerate(structs):
    rr = top + 2 + i
    ce.cell(row=rr, column=1, value=name).font = BLACK
    for j, f in enumerate(cash):
        c = ce.cell(row=rr, column=2 + j, value=f); c.number_format = CUR; c.font = BLACK
    c = ce.cell(row=rr, column=5, value=f"=SUM(B{rr}:D{rr})"); c.number_format = CUR; c.font = BLACK
    c = ce.cell(row=rr, column=6, value=eq if eq != "0" else "=0"); c.number_format = CUR; c.font = BLACK
    c = ce.cell(row=rr, column=7, value=f"=E{rr}+F{rr}"); c.number_format = CUR; c.font = BLACK
    c = ce.cell(row=rr, column=8, value=f"=IF(SUM({rev[0]},{rev[1]},{rev[2]})=0,0,E{rr}/SUM({rev[0]},{rev[1]},{rev[2]}))"); c.number_format = PCT; c.font = BLACK
note = top + 2 + len(structs) + 1
ce.cell(row=note, column=1, value="Reading this table: cash-heavy structures front-load risk onto the company before demand is proven; royalty and equity structures align incentives but dilute margin or ownership. A mixed structure with a modest retainer, a capped royalty, and vesting equity tied to deliverables (content days, launch obligations) is the analyst's recommended starting point, subject to negotiation and counsel. Independent of structure: medical decisions must be insulated from partner influence, and disclosure of ownership/compensation is required in endorsements.").font = Font(name=FONT, italic=True, size=9)
ce.column_dimensions["A"].width = 64
for j in range(2, 9): ce.column_dimensions[L(j)].width = 20

# ------------------------------------------------------------------ Assumptions register (mirror with links)
ar = wb.create_sheet("Assumptions Register")
ar["A1"] = "Assumptions register (links to the Assumptions sheet; edit values there)"; ar["A1"].font = H1
for i, h in enumerate(["Assumption", "Unit", "Base value", "Rationale / source", "Confidence", "Sensitivity", "Validation method"], 1):
    c = ar.cell(row=3, column=i, value=h); c.font = H2; c.fill = HDR
rr = 4
for item in A:
    if item[0] == "SECTION": continue
    key, label, unit, cons, base, up, rat, conf, sens, val = item
    ar.cell(row=rr, column=1, value=label).font = BLACK
    ar.cell(row=rr, column=2, value=unit).font = BLACK
    c = ar.cell(row=rr, column=3, value=f"=Assumptions!$E${ROW[key]}"); c.font = GREEN
    c.number_format = PCT if unit == "%" else (CUR if unit == "$" else '0.00')
    ar.cell(row=rr, column=4, value=rat).font = Font(name=FONT, size=9)
    ar.cell(row=rr, column=5, value=conf).font = BLACK
    ar.cell(row=rr, column=6, value=sens).font = BLACK
    ar.cell(row=rr, column=7, value=val).font = Font(name=FONT, size=9)
    rr += 1
for col, w in zip("ABCDEFG", (62, 9, 14, 80, 11, 11, 26)): ar.column_dimensions[col].width = w

# ------------------------------------------------------------------ README
rd = wb.create_sheet("README", 0)
rd["A1"] = "AminoLord — Driver-based financial model (proposed concept; prepared 2026-09-02)"; rd["A1"].font = H1
lines = [
 "Purpose: illustrate the economics of a compliant premium consumer-health + telehealth + commerce model under three scenarios. Nothing here is a forecast of results; every number is an assumption or a formula on assumptions.",
 "How to use: edit BLUE cells on the Assumptions sheet only. Yellow-filled cells are the highest-sensitivity inputs and should be validated first (see Validation method column).",
 "Sheets: Scenario Summary (outputs) · Assumptions (inputs) · Assumptions Register (rationale) · Model_Conservative / Model_Base / Model_Upside (36 monthly columns) · Sensitivity (LTV:CAC, payback, price×churn grids) · Celebrity Economics (hypothetical partner structures).",
 "Color code: blue = input; black = formula; green = link to another sheet; yellow fill = key assumption.",
 "Structure caveats: prescription-program revenue is modeled as an all-in monthly program price. Whether the brand entity, the affiliated medical practice, or the pharmacy is merchant of record for medication depends on corporate-practice-of-medicine and fee-splitting analysis by counsel; revenue recognition may differ (e.g., management-fee model). Clinician prescribing rate is a clinical outcome that is monitored for safety, never targeted.",
 "Scenario A vs B: set 'Named-partner model active?' to 1 (Scenario A: brand-partner retainer, royalty and traffic uplift apply) or 0 (Scenario B: independent brand).",
 "Sources for benchmark-informed assumptions are summarized in 08_AminoLord_Evidence_Ledger.xlsx and research/06_market_size_benchmarks.md; each assumption's rationale is on the Assumptions sheet.",
 "This model is not investment, legal, medical, or tax advice.",
]
for i, t in enumerate(lines):
    c = rd.cell(row=3 + i, column=1, value=t); c.font = BLACK; c.alignment = Alignment(wrap_text=True, vertical="top")
    rd.row_dimensions[3 + i].height = 45
rd.column_dimensions["A"].width = 140

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved", OUT, "assumption rows", len(ROW))
